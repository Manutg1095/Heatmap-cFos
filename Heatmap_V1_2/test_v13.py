import ast
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import pandas as pd
from openpyxl import load_workbook

import app
import batch_project as batch
import cfos_core as core
import iniciar


ROOT = Path(__file__).resolve().parent


def catalogs():
    groups = pd.DataFrame(
        [{"Experimento": "HB_04", "Animal": "R1", "Grupo": "SCOP-HB"}]
    )
    bregmas = pd.DataFrame(
        [
            {
                "Experimento": "HB_04",
                "Animal": "R1",
                "Corte": "C5",
                "Region": "BLAa",
                "Bregma_mm": -2.30,
            }
        ]
    )
    return groups, bregmas


def create_named_inputs(root):
    folders = {}
    for label in ("objetos", "area", "roi"):
        folders[label] = root / label
        folders[label].mkdir()
    stem = "HB_04_R1_C5_BLAa_20X"
    objects = folders["objetos"] / f"Statistics_{stem}.xlsx"
    area = folders["area"] / f"Statistics_{stem}_Area.xlsx"
    roi = folders["roi"] / f"ROISet_{stem}.roi"
    for path in (objects, area, roi):
        path.touch()
    return folders, objects, area, roi


class Version13Tests(unittest.TestCase):
    def test_version_splash_and_installer_contract(self):
        self.assertEqual(core.APP_VERSION, "1.2.4")
        self.assertEqual(app.PRODUCT_NAME, "Heatmap V1.2")
        self.assertGreaterEqual(app.SPLASH_MIN_MS, 7000)
        self.assertEqual(
            app.OWNERSHIP_TEXT,
            "Propiedad del Laboratorio de neurobiología conductual VMTG UNAM",
        )
        self.assertTrue(app.SPLASH_IMAGE.exists())
        self.assertIn("openpyxl", iniciar.DEPENDENCIES)
        source = (ROOT / "app.py").read_text(encoding="utf-8")
        self.assertIn("SPLASH_MIN_MS - elapsed_ms", source)

    def test_application_contains_batch_workflow_actions(self):
        tree = ast.parse((ROOT / "app.py").read_text(encoding="utf-8"))
        methods = {
            node.name
            for node in ast.walk(tree)
            if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef))
        }
        for expected in (
            "scan", "save_current_experiment", "load_saved_experiment", "run",
            "open_roi_orientation_png", "mark_all_rois_reviewed",
            "open_initial_roi_png", "toggle_mounting_cell", "select_existing_project",
        ):
            self.assertIn(expected, methods)

    def test_project_path_recovers_single_legacy_sibling(self):
        with tempfile.TemporaryDirectory() as directory:
            parent = Path(directory)
            new_app = parent / "Heatmap_V1_1_nueva"
            new_app.mkdir()
            old_project = parent / "Heatmap_V1_1_anterior" / "proyecto_cfos"
            old_project.mkdir(parents=True)
            (old_project / "registro_general_stacks.xlsx").touch()
            pointer = parent / "pointer_inexistente.txt"
            resolved = app.resolve_project_path(new_app, pointer=pointer)
            self.assertEqual(resolved, old_project.resolve())

    def test_project_path_defaults_outside_application_folder(self):
        with tempfile.TemporaryDirectory() as directory:
            parent = Path(directory)
            app_root = parent / "Heatmap_V1_1"
            app_root.mkdir()
            resolved = app.resolve_project_path(app_root, pointer=parent / "none.txt")
            self.assertEqual(resolved, (parent / "proyecto_cfos").resolve())

    def test_mounting_rule_derives_only_x_or_z(self):
        record = {
            "Hemisferio": "L",
            "Direccion_Z_AP": "Posterior",
            "Flip_X": "No",
            "Flip_Y": "No",
            "Revision_ROI": "Revisado",
            "QC": "Excluir",
        }
        left = batch.set_mounting(record, "No")
        self.assertEqual(left["Hemisferio"], "L")
        self.assertEqual(left["Flip_X"], "Si")
        self.assertEqual(left["Flip_Z"], "No")
        self.assertEqual(left["Flip_Y"], "No")
        self.assertEqual(left["Revision_ROI"], "Pendiente")
        right = batch.set_mounting(record, "Si")
        self.assertEqual(right["Flip_X"], "No")
        self.assertEqual(right["Flip_Z"], "Si")
        self.assertEqual(right["Direccion_Z_AP"], "Posterior")

    def test_scan_uses_internal_standard_orientation_without_user_setting(self):
        groups, bregmas = catalogs()
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            folders, _, _, _ = create_named_inputs(root)
            with patch.object(
                batch, "load_area", return_value={"Area_um2": 500_000, "Volumen_um3": 10_000_000}
            ), patch.object(
                batch, "_preflight_counts", return_value={"Objetos_XYZ": 1, "Objetos_incluidos_filtro": 1, "Objetos_validos": 1}
            ):
                staging, _ = batch.scan_experiment_sets(
                    folders["objetos"], folders["area"], folders["roi"], groups, bregmas
                )
        self.assertEqual(staging.iloc[0]["Hemisferio"], "L")
        self.assertEqual(staging.iloc[0]["Direccion_Z_AP"], batch.STANDARD_Z_DIRECTION)
        source = (ROOT / "app.py").read_text(encoding="utf-8")
        self.assertNotIn("Z aumenta hacia", source)
        self.assertIn("Montado de lado derecho", source)

    def test_preregistration_uses_black_and_white_rows(self):
        source = (ROOT / "app.py").read_text(encoding="utf-8")
        self.assertIn('background="#FFFFFF", foreground="#000000"', source)
        self.assertNotIn('background="#FFF6D9"', source)

    def test_scan_pairs_three_folders_and_builds_auditable_row(self):
        groups, bregmas = catalogs()
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            folders, objects, area, roi = create_named_inputs(root)
            counts = {
                "Objetos_XYZ": 100,
                "Objetos_incluidos_filtro": 82,
                "Objetos_fuera_ROI": 7,
                "Objetos_validos": 75,
                "Alineacion_XY": "ROI_e_Imaris_mismo_origen",
            }
            with patch.object(
                batch, "load_area", return_value={"Area_um2": 500_000, "Volumen_um3": 10_000_000}
            ), patch.object(batch, "_preflight_counts", return_value=counts):
                staging, report = batch.scan_experiment_sets(
                    folders["objetos"], folders["area"], folders["roi"], groups, bregmas
                )
        self.assertEqual(report["Experimento"], "HB_04")
        self.assertEqual(len(staging), 1)
        row = staging.iloc[0]
        self.assertEqual(row["Stack_ID"], "HB_04_R1_C5_BLAa")
        self.assertEqual(row["Grupo"], "SCOP-HB")
        self.assertAlmostEqual(float(row["Bregma_mm"]), -2.30)
        self.assertEqual(int(row["Objetos_validos"]), 75)
        self.assertEqual(row["QC"], "Revisar")
        self.assertNotIn("dirección Z", row["Detalle_QC"])
        self.assertEqual(row["Revision_ROI"], "Pendiente")
        self.assertIn("ROI no revisada", row["Detalle_QC"])
        self.assertEqual(row["Montado_lado_derecho"], "No")

        updated = row.to_dict()
        updated = batch.set_mounting(updated, "No")
        updated["Revision_ROI"] = "Revisado"
        updated = batch.refresh_qc(updated)
        # Al salir del directorio temporal las entradas ya no existen; la nueva
        # validación debe impedir que el registro parezca ejecutable.
        self.assertEqual(updated["QC"], "Revisar")
        self.assertIn("No existe Archivo_objetos", updated["Detalle_QC"])

    def test_roi_orientation_png_contains_all_stacks(self):
        groups, bregmas = catalogs()
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            folders, objects, area, roi = create_named_inputs(root)
            counts = {
                "Objetos_XYZ": 100,
                "Objetos_incluidos_filtro": 82,
                "Objetos_fuera_ROI": 7,
                "Objetos_validos": 75,
                "Alineacion_XY": "ROI_e_Imaris_mismo_origen",
            }
            import numpy as np
            polygon = np.array([[10, 10], [90, 20], [75, 80], [25, 95]], dtype=float)
            with patch.object(
                batch, "load_area", return_value={"Area_um2": 500_000, "Volumen_um3": 10_000_000}
            ), patch.object(batch, "_preflight_counts", return_value=counts):
                staging, _ = batch.scan_experiment_sets(
                    folders["objetos"], folders["area"], folders["roi"], groups, bregmas
                )
            with patch.object(batch, "read_roi_polygons", return_value=[polygon]):
                output, report = batch.generate_roi_orientation_png(root / "project", staging, apply_corrections=False)
            self.assertTrue(output.is_file())
            self.assertGreater(output.stat().st_size, 5_000)
            self.assertEqual(report["ROIs_visibles"], 1)
            self.assertEqual(report["Total"], 1)
            self.assertEqual(report["Errores"], [])

    def test_duplicate_candidate_is_flagged_instead_of_guessed(self):
        groups, bregmas = catalogs()
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            folders, objects, area, roi = create_named_inputs(root)
            (folders["objetos"] / "Copy_HB_04_R1_C5_BLAa.xlsx").touch()
            with patch.object(
                batch, "load_area", return_value={"Area_um2": 500_000, "Volumen_um3": 10_000_000}
            ):
                staging, _ = batch.scan_experiment_sets(
                    folders["objetos"], folders["area"], folders["roi"], groups, bregmas
                )
        row = staging.iloc[0]
        self.assertEqual(row["QC"], "Revisar")
        self.assertIn("2 candidatos para Archivo_objetos", row["Detalle_QC"])

    def test_save_replaces_registry_without_duplicate_and_preserves_versions(self):
        groups, bregmas = catalogs()
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            folders, objects, area, roi = create_named_inputs(root)
            counts = {
                "Objetos_XYZ": 100,
                "Objetos_incluidos_filtro": 82,
                "Objetos_fuera_ROI": 7,
                "Objetos_validos": 75,
                "Alineacion_XY": "ROI_e_Imaris_mismo_origen",
            }
            with patch.object(
                batch, "load_area", return_value={"Area_um2": 500_000, "Volumen_um3": 10_000_000}
            ), patch.object(batch, "_preflight_counts", return_value=counts):
                staging, _ = batch.scan_experiment_sets(
                    folders["objetos"], folders["area"], folders["roi"], groups, bregmas
                )
            staging.loc[:, "Direccion_Z_AP"] = "Posterior"
            staging.loc[:, "Revision_ROI"] = "Revisado"
            staging = pd.DataFrame(
                [batch.refresh_qc(core.apply_mounting_correction(row.to_dict())) for _, row in staging.iterrows()],
                columns=core.META_COLUMNS,
            )
            project = root / "project"
            first = batch.save_experiment(project, staging)
            self.assertEqual(len(first["Data"]), 1)
            with self.assertRaises(batch.DuplicateStackError):
                batch.save_experiment(project, staging)
            second = batch.save_experiment(project, staging, overwrite=True)
            self.assertEqual(len(second["Data"]), 1)
            versions = list(
                (project / "experimentos" / "HB_04" / "entradas" / "HB_04_R1_C5_BLAa").iterdir()
            )
            self.assertEqual(len(versions), 2)

            workbook = load_workbook(project / batch.COUNTS_NAME, read_only=True, data_only=False)
            self.assertEqual(
                workbook.sheetnames,
                [
                    "Conteo_por_stack",
                    "Conteo_por_animal",
                    "Resumen_grupo_region",
                    "Cobertura_bregma",
                    "Resumen_experimentos",
                ],
            )
            counts_sheet = workbook["Resumen_grupo_region"]
            headers = [cell.value for cell in counts_sheet[1]]
            values = [cell.value for cell in counts_sheet[2]]
            summary = dict(zip(headers, values))
            self.assertEqual(summary["Grupo"], "SCOP-HB")
            self.assertEqual(summary["Region"], "BLAa")
            self.assertEqual(summary["Objetos_validos_total"], 75)
            workbook.close()

    def test_project_owns_inputs_and_repairs_paths_after_move(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "fuentes_originales"
            source.mkdir()
            objects = source / "Statistics_HB_04_R1_C5_BLAa.xlsx"
            area = source / "Statistics_HB_04_R1_C5_BLAa_Area.xlsx"
            roi = source / "ROISet_HB_04_R1_C5_BLAa.roi"
            for path in (objects, area, roi):
                path.touch()

            record = {column: "" for column in core.META_COLUMNS}
            record.update(
                {
                    "Stack_ID": "HB_04_R1_C5_BLAa", "Experimento": "HB_04",
                    "Animal": "R1", "Corte": "C5", "Region": "BLAa",
                    "Grupo": "SCOP-HB", "Bregma_mm": -2.3,
                    "Montado_lado_derecho": "No", "Revision_ROI": "Revisado",
                    "QC": "OK", "Usar_mapa_grupal": "Si", "Area_um2": 500_000,
                    "Volumen_um3": 10_000_000, "Objetos_validos": 75,
                    "Archivo_objetos": str(objects), "Archivo_area": str(area),
                    "Archivo_ROI": str(roi),
                }
            )
            legacy_app = root / "Heatmap_V1_1_anterior"
            project = legacy_app / "proyecto_cfos"
            batch.save_experiment(project, pd.DataFrame([record], columns=core.META_COLUMNS))

            workbook = load_workbook(project / batch.REGISTRY_NAME)
            sheet = workbook["Registro_stacks"]
            headers = {cell.value: cell.column for cell in sheet[1]}
            relative_paths = {}
            for key in batch.FILE_KEYS:
                value = sheet.cell(2, headers[key]).value
                self.assertFalse(Path(value).is_absolute())
                self.assertTrue(value.startswith("experimentos/"))
                relative_paths[key] = value
                # Simula exactamente un registro viejo con rutas absolutas.
                sheet.cell(2, headers[key]).value = str((project / value).resolve())
            workbook.save(project / batch.REGISTRY_NAME)
            workbook.close()

            for path in (objects, area, roi):
                path.unlink()
            stable = root / "proyecto_cfos"
            project.rename(stable)

            loaded = batch.load_registry(stable)
            self.assertEqual(loaded.iloc[0]["QC"], "OK")
            for key in batch.FILE_KEYS:
                resolved = Path(loaded.iloc[0][key])
                self.assertTrue(resolved.is_file())
                self.assertTrue(resolved.is_relative_to(stable))
            self.assertTrue((stable / "REPORTE_REPARACION_RUTAS.txt").is_file())

            repaired = load_workbook(stable / batch.REGISTRY_NAME, read_only=True, data_only=True)
            repaired_sheet = repaired["Registro_stacks"]
            repaired_headers = {cell.value: cell.column for cell in repaired_sheet[1]}
            for key in batch.FILE_KEYS:
                self.assertEqual(repaired_sheet.cell(2, repaired_headers[key]).value, relative_paths[key])
            repaired.close()

    def test_pointer_prefers_stable_project_over_old_app(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory) / "Heatmap_V1_1_nueva"
            root.mkdir()
            old = root.parent / "Heatmap_V1_1_anterior" / "proyecto_cfos"
            stable = root.parent / "proyecto_cfos"
            old.mkdir(parents=True)
            stable.mkdir()
            (old / "registro_general_stacks.xlsx").touch()
            (stable / "registro_general_stacks.xlsx").touch()
            pointer = root.parent / "project_path.txt"
            pointer.write_text(str(old), encoding="utf-8")
            self.assertEqual(app.resolve_project_path(root, pointer), stable.resolve())

    def test_master_project_overrides_any_saved_pointer(self):
        with tempfile.TemporaryDirectory() as directory:
            package = Path(directory) / "Heatmap_cFos_V1_2"
            root = package / "Heatmap_V1_2"
            stable = package / "proyecto_cfos"
            old = Path(directory) / "Heatmap_V1_1" / "proyecto_cfos"
            root.mkdir(parents=True)
            stable.mkdir()
            old.mkdir(parents=True)
            (stable / "registro_general_stacks.xlsx").touch()
            (stable / app.MASTER_PROJECT_MARKER).touch()
            (old / "registro_general_stacks.xlsx").touch()
            pointer = Path(directory) / "project_path.txt"
            pointer.write_text(str(old), encoding="utf-8")
            self.assertEqual(app.resolve_project_path(root, pointer), stable.resolve())

    def test_master_experiments_are_editable(self):
        self.assertEqual(app.FROZEN_EXPERIMENTS, set())

    def test_bundled_catalogs_start_empty_and_valid(self):
        groups = core.load_group_catalog(ROOT / "catalogos" / "asignacion_grupos.xlsx")
        bregmas = core.load_bregma_catalog(ROOT / "catalogos" / "registro_bregma.xlsx")
        self.assertEqual(len(groups), 0)
        self.assertEqual(len(bregmas), 0)
        joined = bregmas.merge(groups, on=["Experimento", "Animal"], how="left")
        self.assertTrue(joined["Grupo"].notna().all())


if __name__ == "__main__":
    unittest.main(verbosity=2)
