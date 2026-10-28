"""Gestión por experimentos, pre-registro y conteos para Heatmap V1.2."""
from __future__ import annotations

import math
import shutil
from datetime import datetime
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from cfos_core import (
    META_COLUMNS,
    RANGES,
    apply_mounting_correction,
    align_cells_and_roi,
    load_area,
    load_cells,
    numeric_value,
    normalized_label,
    orient_xy,
    parse_name,
    read_roi_polygons,
    portable_project_file_path,
    resolve_project_file_path,
    resolve_record_file_paths,
    save_figure_atomic,
    safe_name,
    selected_for_analysis,
    text,
    to_anatomical_y,
    truthy,
    yes_no,
)


REGISTRY_NAME = "registro_general_stacks.xlsx"
COUNTS_NAME = "conteo_objetos.xlsx"
FILE_KEYS = ("Archivo_objetos", "Archivo_area", "Archivo_ROI")
QC_REQUIRED = (
    "Stack_ID",
    "Grupo",
    "Bregma_mm",
    "Archivo_objetos",
    "Archivo_area",
    "Archivo_ROI",
    "Area_um2",
    "Volumen_um3",
)

STANDARD_Z_DIRECTION = "Posterior"


class DuplicateStackError(RuntimeError):
    def __init__(self, stack_ids):
        self.stack_ids = sorted(set(stack_ids))
        super().__init__("Ya existen estos Stack_ID: " + ", ".join(self.stack_ids))


def blank_record():
    record = {column: "" for column in META_COLUMNS}
    record["Incluir_analisis"] = "Si"
    return record


def stack_id_key(value):
    """Clave canónica de Stack_ID para impedir duplicados por formato o índice."""
    raw = text(value)
    if not raw:
        return ""
    try:
        return parse_name(raw)["Stack_ID"].casefold()
    except ValueError:
        return raw.casefold()


def _record_completeness_score(project, record):
    existing_files = 0
    declared_files = 0
    for key in FILE_KEYS:
        if text(record.get(key)):
            declared_files += 1
        if resolve_project_file_path(project, record.get(key), record, key).is_file():
            existing_files += 1
    populated = sum(
        bool(text(record.get(key)))
        for key in (
            "Grupo", "Bregma_mm", "Area_um2", "Volumen_um3",
            "Objetos_XYZ", "Objetos_validos", "Revision_ROI",
        )
    )
    qc_ok = int(text(record.get("QC")).casefold() == "ok")
    return existing_files, declared_files, populated, qc_ok, text(record.get("Actualizado"))


def consolidate_duplicate_stacks(project, frame):
    """Conserva una sola fila completa por Stack_ID y fusiona datos no vacíos.

    La fila con más archivos existentes (objetos, área y ROI) es la base. Una
    corrección manual de bregma se conserva, pero una fila incompleta nunca
    reemplaza los vínculos válidos del stack.
    """
    project = Path(project)
    data = frame.reindex(columns=META_COLUMNS, fill_value="").copy().reset_index(drop=True)
    if data.empty:
        return data, {"Stack_ID": [], "Filas_eliminadas": 0}

    buckets = {}
    order = []
    for position, row in data.iterrows():
        key = stack_id_key(row.get("Stack_ID")) or f"__fila_sin_id_{position}"
        if key not in buckets:
            buckets[key] = []
            order.append(key)
        buckets[key].append(row.to_dict())

    consolidated = []
    duplicate_ids = []
    removed = 0
    for key in order:
        records = buckets[key]
        if len(records) == 1:
            consolidated.append(records[0])
            continue

        removed += len(records) - 1
        display_id = text(records[0].get("Stack_ID"))
        try:
            display_id = parse_name(display_id)["Stack_ID"]
        except ValueError:
            pass
        duplicate_ids.append(display_id)

        base = dict(max(records, key=lambda item: _record_completeness_score(project, item)))
        for candidate in records:
            for column in META_COLUMNS:
                if not text(base.get(column)) and text(candidate.get(column)):
                    base[column] = candidate.get(column)

        manual = [
            item for item in records
            if normalized_label(item.get("Fuente_bregma")) == "edicionmanual"
            and numeric_value(item.get("Bregma_mm")) is not None
        ]
        if manual:
            chosen = max(manual, key=lambda item: text(item.get("Actualizado")))
            base["Bregma_mm"] = chosen["Bregma_mm"]
            base["Fuente_bregma"] = "Edición manual"
        latest = max((text(item.get("Actualizado")) for item in records), default="")
        if latest:
            base["Actualizado"] = latest
        base["Stack_ID"] = display_id
        consolidated.append(refresh_qc(base))

    result = pd.DataFrame(consolidated, columns=META_COLUMNS).reset_index(drop=True)
    return result, {"Stack_ID": duplicate_ids, "Filas_eliminadas": removed}


def replace_staging_row(staging, position, record):
    """Reemplaza una fila por posición visual sin crear una etiqueta nueva."""
    data = staging.reindex(columns=META_COLUMNS, fill_value="").copy().reset_index(drop=True)
    if position < 0 or position >= len(data):
        raise IndexError("La fila seleccionada ya no existe en el pre-registro.")
    data.iloc[position, :] = [record.get(column, "") for column in META_COLUMNS]
    return data


def resolve_stack_metadata(parsed, groups, bregmas, area_info=None, prefer_catalog=False):
    """Completa grupo y bregma con una prioridad explícita y auditable.

    En el escaneo inicial se conserva la compatibilidad histórica: el Excel de
    área tiene prioridad. Al actualizar el catálogo, ``prefer_catalog=True``
    permite aplicar el nuevo valor sin perder ``Bregma_area_mm`` como dato de
    auditoría.
    """
    area_info = area_info or {}
    group_hit = groups[
        groups["Experimento"].str.casefold().eq(parsed["Experimento"].casefold())
        & groups["Animal"].str.casefold().eq(parsed["Animal"].casefold())
    ]
    group = "" if len(group_hit) != 1 else str(group_hit.iloc[0]["Grupo"])

    bregma_hit = bregmas[
        bregmas["Experimento"].str.casefold().eq(parsed["Experimento"].casefold())
        & bregmas["Animal"].str.casefold().eq(parsed["Animal"].casefold())
        & bregmas["Corte"].str.casefold().eq(parsed["Corte"].casefold())
        & bregmas["Region"].str.casefold().eq(parsed["Region"].casefold())
    ]
    catalog_bregma = None if len(bregma_hit) != 1 else float(bregma_hit.iloc[0]["Bregma_mm"])
    area_bregma = numeric_value(area_info.get("Bregma_area_mm"))
    if prefer_catalog and catalog_bregma is not None:
        bregma = catalog_bregma
        source = "Registro maestro actualizado"
    elif area_bregma is not None:
        bregma = area_bregma
        source = "Excel de área"
    else:
        bregma = catalog_bregma
        source = "Registro maestro" if catalog_bregma is not None else ""
    return {
        "Grupo": group,
        "Fuente_grupo": "Asignación maestra" if group else "",
        "Bregma_mm": bregma,
        "Fuente_bregma": source,
    }


def _files_by_stack(folder, suffixes):
    folder = Path(folder)
    if not folder.is_dir():
        raise ValueError(f"No existe la carpeta: {folder}")
    matches = {}
    ignored = []
    for path in sorted(p for p in folder.rglob("*") if p.is_file()):
        if path.name.startswith(".") or path.suffix.lower() not in suffixes:
            continue
        try:
            stack_id = parse_name(path)["Stack_ID"]
        except ValueError:
            ignored.append(str(path))
            continue
        matches.setdefault(stack_id, []).append(path.resolve())
    return matches, ignored


def _included_filter_count(cells):
    return int(
        cells["Filtro_aplicado"]
        .astype(str)
        .str.strip()
        .str.casefold()
        .isin(["include", "incluir"])
        .sum()
    )


def _preflight_counts(objects_path, roi_path):
    cells = load_cells(objects_path)
    included = _included_filter_count(cells)
    output = {
        "Objetos_XYZ": int(len(cells)),
        "Objetos_incluidos_filtro": included,
        "Objetos_validos": "",
        "Objetos_fuera_ROI": "",
        "Alineacion_XY": "",
    }
    if not roi_path:
        return output
    try:
        polygons = read_roi_polygons(roi_path)
        aligned, _, alignment, direct, origin = align_cells_and_roi(cells, polygons, "")
        valid = aligned["Filtro_incluye"] & aligned["Dentro_ROI"]
        output.update(
            {
                "Objetos_validos": int(valid.sum()),
                "Objetos_fuera_ROI": int((aligned["Filtro_incluye"] & ~aligned["Dentro_ROI"]).sum()),
                "Alineacion_XY": alignment,
                "Fraccion_dentro_directa": direct,
                "Fraccion_dentro_origen": origin,
            }
        )
    except Exception as exc:
        output["Preflight_error"] = str(exc)
    return output


def quality_issues(record):
    issues = []
    detail = text(record.get("Detalle_QC_base"))
    if detail:
        base_issues = [part.strip() for part in detail.split(" | ") if part.strip()]
        # Los avisos creados durante el escaneo no deben quedar obsoletos si el
        # usuario corrige después el grupo o el bregma.
        if text(record.get("Grupo")):
            base_issues = [
                issue for issue in base_issues
                if normalized_label(issue) not in {"gruponoencontrado", "faltagrupo"}
            ]
        if numeric_value(record.get("Bregma_mm")) is not None:
            base_issues = [
                issue for issue in base_issues
                if normalized_label(issue) not in {"bregmanoencontrado", "faltabregmamm"}
            ]
        issues.extend(base_issues)
    for key in QC_REQUIRED:
        value = record.get(key)
        if key in {"Bregma_mm", "Area_um2", "Volumen_um3"}:
            number = numeric_value(value)
            if number is None or (key != "Bregma_mm" and number <= 0):
                issues.append(f"Falta {key}")
        elif not text(value):
            issues.append(f"Falta {key}")
        elif key in FILE_KEYS and not Path(text(value)).expanduser().is_file():
            issues.append(f"No existe {key}")
    if yes_no(record.get("Montado_lado_derecho")) == "":
        issues.append("Clasificar si fue montado de lado derecho")
    if text(record.get("Revision_ROI")).casefold() != "revisado":
        issues.append("ROI no revisada en el PNG de orientación")
    valid_count = numeric_value(record.get("Objetos_validos"))
    if valid_count is None:
        issues.append("Conteo dentro de ROI no calculado")
    elif valid_count <= 0:
        issues.append("Sin objetos válidos dentro de ROI")
    return list(dict.fromkeys(issues))


def refresh_qc(record, preserve_excluded=True):
    result = dict(record)
    issues = quality_issues(result)
    if preserve_excluded and text(result.get("QC")).casefold() == "excluir":
        result["QC"] = "Excluir"
    else:
        result["QC"] = "OK" if not issues else "Revisar"
    result["Detalle_QC"] = " | ".join(issues)
    return result


def set_mounting(record, mounted_right, reset_review=None):
    """Guarda la única decisión manual y deriva Flip X/Flip Z."""
    result = dict(record)
    previous = yes_no(result.get("Montado_lado_derecho"))
    requested = yes_no(mounted_right)
    result["Montado_lado_derecho"] = requested
    result = apply_mounting_correction(result)
    if reset_review is None:
        reset_review = previous != requested
    if reset_review:
        result["Revision_ROI"] = "Pendiente"
    return refresh_qc(result)


def set_bregma(record, value, source="Edición manual"):
    """Valida y aplica un bregma editable sin alterar el valor del Excel de área."""
    result = dict(record)
    bregma = numeric_value(value)
    if bregma is None:
        raise ValueError("Escribe un Bregma numérico en milímetros.")
    region = text(result.get("Region"))
    if region not in RANGES:
        raise ValueError(f"No se puede validar el Bregma: región no reconocida ({region or 'vacía'}).")
    anterior, posterior = RANGES[region]
    if not posterior <= bregma <= anterior:
        raise ValueError(
            f"Bregma {bregma:g} fuera del rango de {region}: "
            f"{posterior:g} a {anterior:g} mm."
        )
    result["Bregma_mm"] = float(bregma)
    result["Fuente_bregma"] = source
    return refresh_qc(result)


def scan_experiment_sets(objects_dir, area_dir, roi_dir, groups, bregmas):
    """Empareja tres carpetas por Stack_ID y devuelve un pre-registro auditable."""
    maps = {}
    ignored = {}
    for key, folder, suffixes in (
        ("Archivo_objetos", objects_dir, {".xlsx", ".xls"}),
        ("Archivo_area", area_dir, {".xlsx", ".xls"}),
        ("Archivo_ROI", roi_dir, {".roi", ".zip"}),
    ):
        maps[key], ignored[key] = _files_by_stack(folder, suffixes)

    stack_ids = sorted(set().union(*(set(mapping) for mapping in maps.values())))
    if not stack_ids:
        raise ValueError("No encontré archivos con nombres de stack reconocibles en las tres carpetas.")
    experiments = {parse_name(stack_id)["Experimento"] for stack_id in stack_ids}
    if len(experiments) != 1:
        raise ValueError(
            "Cada carga debe contener un solo experimento. Encontré: "
            + ", ".join(sorted(experiments))
        )

    rows = []
    for stack_id in stack_ids:
        parsed = parse_name(stack_id)
        record = blank_record()
        record.update(parsed)
        record.update(
            {
                "Hemisferio": "L",
                "Montado_lado_derecho": "No",
                "Direccion_Z_AP": STANDARD_Z_DIRECTION,
                "Flip_X": "No",
                "Flip_Y": "No",
                "Flip_Z": "No",
                "Revision_ROI": "Pendiente",
                "Usar_mapa_grupal": "Si",
                "Estado_registro": "Pre-registro",
            }
        )
        issues = []
        selected = {}
        for key in FILE_KEYS:
            candidates = maps[key].get(stack_id, [])
            if len(candidates) == 1:
                selected[key] = str(candidates[0])
                record[key] = str(candidates[0])
            elif not candidates:
                issues.append(f"Falta {key}")
            else:
                issues.append(f"{len(candidates)} candidatos para {key}")

        area_info = {}
        if selected.get("Archivo_area"):
            try:
                area_info = load_area(selected["Archivo_area"])
                for key, value in area_info.items():
                    record[key] = value
            except Exception as exc:
                issues.append(f"Área/volumen: {exc}")

        record.update(resolve_stack_metadata(parsed, groups, bregmas, area_info))
        if not record["Grupo"]:
            issues.append("Grupo no encontrado")
        if numeric_value(record["Bregma_mm"]) is None:
            issues.append("Bregma no encontrado")

        if selected.get("Archivo_objetos"):
            try:
                preflight = _preflight_counts(
                    selected["Archivo_objetos"], selected.get("Archivo_ROI", "")
                )
                preflight_error = preflight.pop("Preflight_error", "")
                record.update(preflight)
                if preflight_error:
                    issues.append(f"Objetos/ROI: {preflight_error}")
            except Exception as exc:
                issues.append(f"Objetos/ROI: {exc}")
        record["Detalle_QC_base"] = " | ".join(issues)
        record = refresh_qc(record)
        rows.append({column: record.get(column, "") for column in META_COLUMNS})

    report = {
        "Experimento": next(iter(experiments)),
        "Stacks": len(rows),
        "Ignorados": sum(len(paths) for paths in ignored.values()),
        "Ignorados_por_tipo": ignored,
    }
    return pd.DataFrame(rows, columns=META_COLUMNS), report


def generate_roi_orientation_png(project, staging, experiment=None, apply_corrections=False):
    """Crea el mosaico inicial crudo o el mosaico corregido del análisis."""
    if staging is None or staging.empty:
        raise ValueError("El pre-registro está vacío; primero carga un experimento.")

    data = staging.reindex(columns=META_COLUMNS, fill_value="").copy()
    detected = sorted({text(value) for value in data["Experimento"] if text(value)})
    if experiment is None:
        if len(detected) != 1:
            raise ValueError("El mosaico debe corresponder a un solo experimento.")
        experiment = detected[0]
    experiment = text(experiment)

    data["_bregma_sort"] = pd.to_numeric(data["Bregma_mm"], errors="coerce")
    data = data.sort_values(
        ["Region", "_bregma_sort", "Animal", "Corte"],
        ascending=[True, False, True, True],
        na_position="last",
    ).reset_index(drop=True)

    count = len(data)
    columns = min(5, max(1, int(math.ceil(math.sqrt(count)))))
    rows = int(math.ceil(count / columns))
    figure, axes = plt.subplots(
        rows,
        columns,
        figsize=(3.25 * columns, 3.75 * rows + 1.15),
        squeeze=False,
        facecolor="#F3F6FA",
    )
    rendered = 0
    errors = []

    for index, axis in enumerate(axes.flat):
        if index >= count:
            axis.set_visible(False)
            continue
        record = data.iloc[index].to_dict()
        if apply_corrections:
            record = apply_mounting_correction(record)
        stack_id = text(record.get("Stack_ID")) or f"Fila {index + 1}"
        region = text(record.get("Region"))
        color = "#157A8A" if region == "BLAa" else "#C66A22"
        axis.set_facecolor("white")
        axis.set_xlim(0, 1)
        axis.set_ylim((0, 1) if apply_corrections else (1, 0))
        axis.set_aspect("equal", adjustable="box")
        axis.set_xticks([0, 1], labels=["M", "L"] if apply_corrections else ["X0", "X1"])
        axis.set_yticks([0, 1], labels=["V", "D"] if apply_corrections else ["D", "V"])
        axis.tick_params(labelsize=8, length=2, colors="#344054")
        for spine in axis.spines.values():
            spine.set_color("#CBD5E1")

        roi_path = Path(text(record.get("Archivo_ROI")))
        try:
            if not roi_path.is_file():
                raise FileNotFoundError("archivo ROI no disponible")
            polygons = read_roi_polygons(roi_path)
            all_points = np.vstack(polygons)
            minimum = all_points.min(axis=0)
            maximum = all_points.max(axis=0)
            span = np.where(maximum > minimum, maximum - minimum, 1.0)
            normalized = [(polygon - minimum) / span for polygon in polygons]
            oriented = [
                orient_xy(
                    polygon,
                    "L",
                    truthy(record.get("Flip_X")) if apply_corrections else False,
                    False,
                )
                for polygon in normalized
            ]
            if apply_corrections:
                oriented = [to_anatomical_y(polygon) for polygon in oriented]
            for polygon in oriented:
                closed = np.vstack([polygon, polygon[0]])
                axis.fill(closed[:, 0], closed[:, 1], color=color, alpha=0.13)
                axis.plot(closed[:, 0], closed[:, 1], color=color, linewidth=2.0)
            rendered += 1
        except Exception as exc:
            errors.append(f"{stack_id}: {exc}")
            axis.text(
                0.5,
                0.5,
                "ROI no visible\n" + str(exc),
                ha="center",
                va="center",
                color="#B42318",
                fontsize=8,
                wrap=True,
            )

        bregma = numeric_value(record.get("Bregma_mm"))
        bregma_text = "? mm" if bregma is None else f"{bregma:.2f} mm"
        mounting = yes_no(record.get("Montado_lado_derecho"))
        if not apply_corrections:
            view_text = "Sin clasificar"
        elif mounting == "Si":
            view_text = "Derecho → corrección Z"
        elif mounting == "No":
            view_text = "Izquierdo → corrección X"
        else:
            view_text = "Pendiente de clasificar"
        axis.set_title(
            f"{stack_id}\n"
            f"{text(record.get('Grupo'))} · {bregma_text}\n"
            f"Montaje: {view_text} · "
            f"{text(record.get('Revision_ROI')) or 'Pendiente'}",
            fontsize=8.2,
            fontweight="bold",
            color="#101828",
            pad=5,
            linespacing=1.15,
        )

    title_mode = "MOSAICO INICIAL — antes de clasificar el montaje" if not apply_corrections else "MOSAICO CORREGIDO — vista aplicada al análisis"
    subtitle = (
        "X/Y originales 0–1"
        if not apply_corrections
        else "X/Y normalizadas"
    )
    figure.suptitle(
        f"Control de ROIs — {experiment}\n{title_mode}\n{subtitle}",
        fontsize=15,
        fontweight="bold",
        color="#0B1F33",
        y=0.995,
    )
    figure.text(
        0.5,
        0.008,
        "BLAa en azul · BLAp en naranja · Izquierdo corrige X · Derecho corrige Z.",
        ha="center",
        va="bottom",
        fontsize=9,
        color="#475467",
    )
    figure.tight_layout(rect=(0.02, 0.06, 0.98, 0.90), h_pad=1.7, w_pad=1.2)

    output_dir = Path(project) / "control_roi" / safe_name(experiment)
    output_dir.mkdir(parents=True, exist_ok=True)
    suffix = "corregido" if apply_corrections else "inicial_sin_clasificar"
    output = output_dir / f"control_ROIs_{suffix}_{safe_name(experiment)}.png"
    save_figure_atomic(
        figure, output, dpi=220, bbox_inches="tight", facecolor=figure.get_facecolor()
    )
    plt.close(figure)
    return output, {"ROIs_visibles": rendered, "Errores": errors, "Total": count}


def _as_number(frame, columns):
    for column in columns:
        if column in frame:
            frame[column] = pd.to_numeric(frame[column], errors="coerce")
    return frame


def load_registry(project):
    project = Path(project)
    registry = project / REGISTRY_NAME
    legacy = project / "metadata_stacks.xlsx"
    source = registry if registry.exists() else legacy
    if not source.exists():
        return pd.DataFrame(columns=META_COLUMNS)
    frame = pd.read_excel(source, sheet_name=0, dtype=str).fillna("")
    frame = frame.reindex(columns=META_COLUMNS, fill_value="")
    frame.loc[
        frame["Incluir_analisis"].astype(str).str.strip().eq(""), "Incluir_analisis"
    ] = "Si"
    repaired_paths = []
    resolved_rows = []
    for _, row in frame.iterrows():
        original = row.to_dict()
        resolved = resolve_record_file_paths(project, original)
        for key in FILE_KEYS:
            old = text(original.get(key))
            new = text(resolved.get(key))
            if old and Path(old).is_absolute() and new and new != old and Path(new).is_file():
                repaired_paths.append(
                    f"{text(original.get('Stack_ID'))} · {key}: {old} -> {new}"
                )
        resolved_rows.append(resolved)
    frame = pd.DataFrame(resolved_rows, columns=META_COLUMNS)
    frame.loc[
        frame["Revision_ROI"].astype(str).str.strip().eq(""), "Revision_ROI"
    ] = "Pendiente"
    if not frame.empty:
        frame = pd.DataFrame(
            [refresh_qc(apply_mounting_correction(row.to_dict())) for _, row in frame.iterrows()],
            columns=META_COLUMNS,
        )
    if repaired_paths:
        write_registry(project, frame)
        (project / "REPORTE_REPARACION_RUTAS.txt").write_text(
            "REPARACIÓN AUTOMÁTICA DE RUTAS — HEATMAP V1.2\n"
            f"Fecha: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n"
            f"Archivos reparados: {len(repaired_paths)}\n\n"
            + "\n".join(repaired_paths)
            + "\n",
            encoding="utf-8",
        )
    return frame


def experiment_summary(registry):
    if registry.empty:
        return pd.DataFrame(
            columns=["Experimento", "N_stacks", "N_animales", "BLAa", "BLAp", "QC_OK", "Revisar", "Excluir"]
        )
    rows = []
    for experiment, members in registry.groupby("Experimento", dropna=False):
        qc = members["QC"].astype(str).str.casefold()
        rows.append(
            {
                "Experimento": experiment,
                "N_stacks": len(members),
                "N_animales": members["Animal"].nunique(),
                "BLAa": int(members["Region"].eq("BLAa").sum()),
                "BLAp": int(members["Region"].eq("BLAp").sum()),
                "QC_OK": int(qc.eq("ok").sum()),
                "Revisar": int(qc.eq("revisar").sum()),
                "Excluir": int(qc.eq("excluir").sum()),
            }
        )
    return pd.DataFrame(rows).sort_values("Experimento").reset_index(drop=True)



def _write_only_valid_counts(path, count_map):
    """Actualiza únicamente la columna Objetos_validos en un registro Excel.

    No reescribe metadatos, ROI, bregma, volumen, orientación ni QC.
    """
    path = Path(path)
    if not path.is_file():
        return 0
    workbook = load_workbook(path)
    try:
        sheet = workbook["Registro_stacks"] if "Registro_stacks" in workbook.sheetnames else workbook[workbook.sheetnames[0]]
        headers = {text(cell.value): cell.column for cell in sheet[1] if text(cell.value)}
        stack_col = headers.get("Stack_ID")
        valid_col = headers.get("Objetos_validos")
        if not stack_col or not valid_col:
            raise ValueError(f"{path.name} no contiene Stack_ID y Objetos_validos.")
        changed = 0
        for row in range(2, sheet.max_row + 1):
            stack_id = text(sheet.cell(row, stack_col).value)
            if stack_id not in count_map:
                continue
            new_value = int(count_map[stack_id])
            cell = sheet.cell(row, valid_col)
            old_value = numeric_value(cell.value)
            if old_value is None or int(old_value) != new_value:
                changed += 1
            cell.value = new_value
        workbook.save(path)
        return changed
    finally:
        workbook.close()


def _write_analysis_selection(path, experiment, animal, include):
    """Actualiza únicamente la selección de análisis de un animal."""
    path = Path(path)
    if not path.is_file():
        return 0
    workbook = load_workbook(path)
    try:
        sheet = workbook["Registro_stacks"] if "Registro_stacks" in workbook.sheetnames else workbook[workbook.sheetnames[0]]
        headers = {text(cell.value): cell.column for cell in sheet[1] if text(cell.value)}
        experiment_col = headers.get("Experimento")
        animal_col = headers.get("Animal")
        if not experiment_col or not animal_col:
            raise ValueError(f"{path.name} no contiene Experimento y Animal.")
        selection_col = headers.get("Incluir_analisis")
        if not selection_col:
            selection_col = sheet.max_column + 1
            sheet.cell(1, selection_col).value = "Incluir_analisis"
        value = "Si" if include else "No"
        changed = 0
        for row in range(2, sheet.max_row + 1):
            row_experiment = text(sheet.cell(row, experiment_col).value)
            row_animal = text(sheet.cell(row, animal_col).value)
            if row_experiment != text(experiment) or row_animal != text(animal):
                continue
            cell = sheet.cell(row, selection_col)
            if text(cell.value) != value:
                changed += 1
            cell.value = value
        workbook.save(path)
        return changed
    finally:
        workbook.close()


def set_animal_analysis_selection(project, experiment, animal, include):
    """Persiste la casilla de análisis sin modificar ROI, QC ni metadatos."""
    project = Path(project)
    registry = project / REGISTRY_NAME
    legacy = project / "metadata_stacks.xlsx"
    source = registry if registry.is_file() else legacy
    if not source.is_file():
        raise ValueError("No hay un registro maestro guardado.")
    changed = _write_analysis_selection(source, experiment, animal, include)
    experiment_file = (
        project / "experimentos" / safe_name(experiment)
        / f"registro_{safe_name(experiment)}.xlsx"
    )
    _write_analysis_selection(experiment_file, experiment, animal, include)
    refreshed = load_registry(project)
    return {"Data": refreshed, "Stacks_actualizados": changed}


def refresh_saved_valid_counts(project):
    """Recalcula sólo Objetos_validos de stacks ya guardados.

    El conteo usa los archivos de objetos y ROI ya vinculados al Stack_ID, el
    filtro morfológico actual y los Flip_X/Flip_Y ya guardados. No recalcula ni
    sobrescribe bregma, grupo, área, volumen, orientación, QC u otros campos.
    """
    project = Path(project)
    registry_path = project / REGISTRY_NAME
    legacy = project / "metadata_stacks.xlsx"
    source = registry_path if registry_path.is_file() else legacy
    if not source.is_file():
        raise ValueError("No hay un registro maestro guardado para actualizar.")
    # Leer el registro crudo evita recalcular o normalizar cualquier otro campo.
    registry = pd.read_excel(source, sheet_name=0, dtype=str).fillna("")
    registry = registry.reindex(columns=META_COLUMNS, fill_value="")
    if registry.empty:
        raise ValueError("No hay stacks guardados para actualizar.")

    count_map = {}
    audit_rows = []
    errors = []
    for _, row in registry.iterrows():
        stack_id = text(row.get("Stack_ID"))
        if not stack_id:
            continue
        try:
            record = resolve_record_file_paths(project, row.to_dict())
            cells = load_cells(record["Archivo_objetos"])
            polygons = read_roi_polygons(record["Archivo_ROI"])
            aligned, _, _, _, _ = align_cells_and_roi(
                cells,
                polygons,
                record.get("Hemisferio", ""),
                truthy(record.get("Flip_X")),
                truthy(record.get("Flip_Y")),
            )
            valid_count = int((aligned["Filtro_incluye"] & aligned["Dentro_ROI"]).sum())
            old_value = numeric_value(row.get("Objetos_validos"))
            count_map[stack_id] = valid_count
            audit_rows.append({
                "Stack_ID": stack_id,
                "Objetos_validos_antes": old_value,
                "Objetos_validos_nuevo": valid_count,
                "Diferencia": None if old_value is None else valid_count - int(old_value),
            })
        except Exception as exc:
            errors.append({"Stack_ID": stack_id, "Error": str(exc)})

    if not count_map:
        details = "\n".join(f"{item['Stack_ID']}: {item['Error']}" for item in errors[:8])
        raise RuntimeError("No se pudo recalcular ningún conteo válido.\n" + details)

    changed = _write_only_valid_counts(source, count_map)

    # Mantiene sincronizados los registros por experimento, tocando sólo la misma columna.
    for experiment in sorted({text(value) for value in registry["Experimento"] if text(value)}):
        exp_file = project / "experimentos" / safe_name(experiment) / f"registro_{safe_name(experiment)}.xlsx"
        _write_only_valid_counts(exp_file, count_map)

    # Regenerar sólo el libro derivado de conteos a partir del registro ya actualizado.
    refreshed = pd.read_excel(source, sheet_name=0, dtype=str).fillna("")
    refreshed = refreshed.reindex(columns=META_COLUMNS, fill_value="")
    counts_path = write_counts_workbook(project, refreshed)
    audit = pd.DataFrame(audit_rows)
    return {
        "Procesados": len(count_map),
        "Cambiados": changed,
        "Errores": errors,
        "Auditoria": audit,
        "Conteos": counts_path,
    }

def _style_workbook(path):
    workbook = load_workbook(path)
    try:
        header_fill = PatternFill("solid", fgColor="0B1F33")
        header_font = Font(color="FFFFFF", bold=True)
        ok_fill = PatternFill("solid", fgColor="D9EAD3")
        review_fill = PatternFill("solid", fgColor="FFF2CC")
        exclude_fill = PatternFill("solid", fgColor="F4CCCC")
        for sheet in workbook.worksheets:
            sheet.sheet_view.showGridLines = False
            sheet.freeze_panes = "A2"
            if sheet.max_row >= 1:
                sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            sheet.row_dimensions[1].height = 32
            headers = {cell.value: cell.column for cell in sheet[1] if cell.value}
            for row in range(2, sheet.max_row + 1):
                for name in ("QC", "Estado_QC"):
                    column = headers.get(name)
                    if not column:
                        continue
                    cell = sheet.cell(row, column)
                    status = text(cell.value).casefold()
                    cell.fill = ok_fill if status == "ok" else exclude_fill if status == "excluir" else review_fill
            for column in range(1, sheet.max_column + 1):
                values = [text(sheet.cell(row, column).value) for row in range(1, min(sheet.max_row, 250) + 1)]
                width = min(42, max(10, max((len(value) for value in values), default=0) + 2))
                sheet.column_dimensions[get_column_letter(column)].width = width
            for name, column in headers.items():
                if name in {"Bregma_mm", "Bregma_min_mm", "Bregma_max_mm"}:
                    for row in range(2, sheet.max_row + 1):
                        sheet.cell(row, column).number_format = "0.00"
                elif "Volumen" in str(name) or "Area" in str(name):
                    for row in range(2, sheet.max_row + 1):
                        sheet.cell(row, column).number_format = "#,##0.00"
                elif "Densidad" in str(name):
                    for row in range(2, sheet.max_row + 1):
                        sheet.cell(row, column).number_format = "#,##0.0"
        workbook.save(path)
    finally:
        workbook.close()


def _write_workbook(path, sheets):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.stem}_tmp.xlsx")
    with pd.ExcelWriter(temporary, engine="openpyxl") as writer:
        for sheet_name, frame in sheets.items():
            frame.to_excel(writer, index=False, sheet_name=sheet_name[:31])
    _style_workbook(temporary)
    temporary.replace(path)
    return path


def portable_registry_frame(project, registry):
    project = Path(project)
    registry = registry.reindex(columns=META_COLUMNS, fill_value="").copy()
    for key in FILE_KEYS:
        registry[key] = registry[key].map(
            lambda value: portable_project_file_path(project, value)
        )
    return registry


def write_registry(project, registry):
    project = Path(project)
    registry = portable_registry_frame(project, registry)
    registry = registry.sort_values(["Experimento", "Animal", "Region", "Corte"]).reset_index(drop=True)
    return _write_workbook(
        project / REGISTRY_NAME,
        {"Registro_stacks": registry, "Resumen_experimentos": experiment_summary(registry)},
    )


def _write_experiment_registry(project, experiment, registry):
    project = Path(project)
    rows = registry[
        registry["Experimento"].astype(str).str.casefold().eq(text(experiment).casefold())
    ].copy()
    path = (
        project / "experimentos" / safe_name(experiment)
        / f"registro_{safe_name(experiment)}.xlsx"
    )
    return _write_workbook(
        path,
        {
            "Registro_stacks": portable_registry_frame(project, rows),
            "Resumen_experimento": experiment_summary(rows),
        },
    )


def _sync_experiment_registries(project, registry):
    outputs = []
    experiments = sorted({text(value) for value in registry["Experimento"] if text(value)})
    for experiment in experiments:
        outputs.append(_write_experiment_registry(project, experiment, registry))
    return outputs


def count_tables(registry):
    registry = registry.reindex(columns=META_COLUMNS, fill_value="").copy()
    numeric_columns = [
        "Bregma_mm", "Area_um2", "Volumen_um3", "Objetos_XYZ",
        "Objetos_incluidos_filtro", "Objetos_fuera_ROI", "Objetos_validos",
    ]
    registry = _as_number(registry, numeric_columns)
    per_stack = registry[
        [
            "Stack_ID", "Experimento", "Animal", "Corte", "Region", "Grupo", "Bregma_mm",
            "Montado_lado_derecho", "Flip_X", "Flip_Z", "QC", "Usar_mapa_grupal",
            "Incluir_analisis",
            "Objetos_XYZ", "Objetos_incluidos_filtro",
            "Objetos_fuera_ROI", "Objetos_validos", "Area_um2", "Volumen_um3",
            "Fuente_grupo", "Fuente_bregma", "Actualizado",
        ]
    ].rename(
        columns={
            "QC": "Estado_QC",
            "Objetos_XYZ": "Objetos_totales_XYZ",
            "Objetos_incluidos_filtro": "Objetos_post_filtro",
            "Objetos_validos": "Objetos_validos_ROI",
        }
    )
    per_stack["Incluido_en_resumen"] = (
        per_stack["Estado_QC"].astype(str).str.casefold().eq("ok")
        & per_stack["Incluir_analisis"].map(selected_for_analysis)
    )
    valid = per_stack[per_stack["Incluido_en_resumen"]].copy()

    animal_columns = [
        "Experimento", "Animal", "Grupo", "Region", "N_stacks",
        "Objetos_totales_XYZ", "Objetos_post_filtro", "Objetos_validos_ROI",
        "Volumen_um3", "Densidad_objetos_validos_mm3", "Bregma_min_mm", "Bregma_max_mm",
    ]
    animal_rows = []
    if not valid.empty:
        for keys, members in valid.groupby(["Experimento", "Animal", "Grupo", "Region"], dropna=False):
            volume = members["Volumen_um3"].sum(min_count=1)
            objects = members["Objetos_validos_ROI"].sum(min_count=1)
            animal_rows.append(
                {
                    "Experimento": keys[0], "Animal": keys[1], "Grupo": keys[2], "Region": keys[3],
                    "N_stacks": len(members),
                    "Objetos_totales_XYZ": members["Objetos_totales_XYZ"].sum(min_count=1),
                    "Objetos_post_filtro": members["Objetos_post_filtro"].sum(min_count=1),
                    "Objetos_validos_ROI": objects,
                    "Volumen_um3": volume,
                    "Densidad_objetos_validos_mm3": objects / volume * 1e9 if pd.notna(volume) and volume > 0 else np.nan,
                    "Bregma_min_mm": members["Bregma_mm"].min(),
                    "Bregma_max_mm": members["Bregma_mm"].max(),
                }
            )
    per_animal = pd.DataFrame(animal_rows, columns=animal_columns)

    group_columns = [
        "Grupo", "Region", "N_animales", "N_stacks", "Objetos_validos_total",
        "Media_objetos_validos_animal", "SD_objetos_validos_animal", "SEM_objetos_validos_animal",
        "Volumen_total_um3", "Densidad_global_mm3",
    ]
    group_rows = []
    if not per_animal.empty:
        for (group, region), members in per_animal.groupby(["Grupo", "Region"], dropna=False):
            n = len(members)
            sd = members["Objetos_validos_ROI"].std(ddof=1) if n > 1 else np.nan
            objects = members["Objetos_validos_ROI"].sum(min_count=1)
            volume = members["Volumen_um3"].sum(min_count=1)
            group_rows.append(
                {
                    "Grupo": group, "Region": region, "N_animales": n,
                    "N_stacks": int(members["N_stacks"].sum()),
                    "Objetos_validos_total": objects,
                    "Media_objetos_validos_animal": members["Objetos_validos_ROI"].mean(),
                    "SD_objetos_validos_animal": sd,
                    "SEM_objetos_validos_animal": sd / math.sqrt(n) if n > 1 else np.nan,
                    "Volumen_total_um3": volume,
                    "Densidad_global_mm3": objects / volume * 1e9 if pd.notna(volume) and volume > 0 else np.nan,
                }
            )
    group_region = pd.DataFrame(group_rows, columns=group_columns)

    coverage_columns = [
        "Grupo", "Region", "Bregma_mm", "N_animales", "N_cortes", "Objetos_validos_ROI", "Volumen_um3",
    ]
    coverage_rows = []
    if not valid.empty:
        for (group, region, bregma), members in valid.groupby(["Grupo", "Region", "Bregma_mm"], dropna=False):
            animals = members[["Experimento", "Animal"]].drop_duplicates()
            coverage_rows.append(
                {
                    "Grupo": group, "Region": region, "Bregma_mm": bregma,
                    "N_animales": len(animals), "N_cortes": len(members),
                    "Objetos_validos_ROI": members["Objetos_validos_ROI"].sum(min_count=1),
                    "Volumen_um3": members["Volumen_um3"].sum(min_count=1),
                }
            )
    coverage = pd.DataFrame(coverage_rows, columns=coverage_columns)
    return {
        "Conteo_por_stack": per_stack,
        "Conteo_por_animal": per_animal,
        "Resumen_grupo_region": group_region,
        "Cobertura_bregma": coverage,
        "Resumen_experimentos": experiment_summary(registry),
    }


def write_counts_workbook(project, registry=None, destination=None):
    project = Path(project)
    registry = load_registry(project) if registry is None else registry
    destination = Path(destination) if destination else project / COUNTS_NAME
    return _write_workbook(destination, count_tables(registry))


def _metadata_key(record):
    return tuple(
        text(record.get(column)).casefold()
        for column in ("Experimento", "Animal", "Corte", "Region")
    )


def update_saved_bregmas_from_catalog(project, bregmas):
    """Sincroniza el catálogo con registros guardados, respetando ediciones manuales."""
    project = Path(project)
    registry, duplicate_report = consolidate_duplicate_stacks(project, load_registry(project))
    if registry.empty:
        return {
            "Data": registry,
            "Coincidencias": 0,
            "Bregmas_cambiados": 0,
            "Registros_actualizados": 0,
            "Omitidos_manual": 0,
            "Sin_coincidencia": 0,
            "Duplicados_reparados": duplicate_report["Filas_eliminadas"],
            "Stack_ID_consolidados": duplicate_report["Stack_ID"],
        }

    required = {"Experimento", "Animal", "Corte", "Region", "Bregma_mm"}
    missing = required - set(bregmas.columns)
    if missing:
        raise ValueError("El catálogo de bregma no contiene: " + ", ".join(sorted(missing)))

    catalog = {}
    for _, row in bregmas.iterrows():
        key = _metadata_key(row)
        if key in catalog:
            raise ValueError("El catálogo contiene una clave de bregma duplicada: " + "/".join(key))
        catalog[key] = float(row["Bregma_mm"])

    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    rows = []
    matched = changed = updated = skipped_manual = unmatched = 0
    for _, row in registry.iterrows():
        record = row.to_dict()
        key = _metadata_key(record)
        if key not in catalog:
            unmatched += 1
            rows.append(record)
            continue
        matched += 1
        if normalized_label(record.get("Fuente_bregma")) == "edicionmanual":
            skipped_manual += 1
            rows.append(record)
            continue

        catalog_value = catalog[key]
        previous = numeric_value(record.get("Bregma_mm"))
        value_changed = previous is None or not math.isclose(
            float(previous), catalog_value, rel_tol=0.0, abs_tol=1e-9
        )
        source_changed = text(record.get("Fuente_bregma")) != "Registro maestro actualizado"
        if value_changed or source_changed:
            record = set_bregma(
                record, catalog_value, source="Registro maestro actualizado"
            )
            if value_changed:
                record["Actualizado"] = now
                changed += 1
            updated += 1
        rows.append(record)

    refreshed = pd.DataFrame(rows, columns=META_COLUMNS)
    if updated or duplicate_report["Filas_eliminadas"]:
        write_registry(project, refreshed)
        _sync_experiment_registries(project, refreshed)
        write_counts_workbook(project, refreshed)
    return {
        "Data": refreshed,
        "Coincidencias": matched,
        "Bregmas_cambiados": changed,
        "Registros_actualizados": updated,
        "Omitidos_manual": skipped_manual,
        "Sin_coincidencia": unmatched,
        "Duplicados_reparados": duplicate_report["Filas_eliminadas"],
        "Stack_ID_consolidados": duplicate_report["Stack_ID"],
    }


def archive_experiment(project, experiment):
    """Retira un experimento activo y conserva una copia recuperable."""
    project = Path(project)
    current = load_registry(project)
    requested = text(experiment)
    if not requested:
        raise ValueError("Selecciona un experimento para eliminar.")
    mask = current["Experimento"].astype(str).str.casefold().eq(requested.casefold())
    removed = current[mask].copy()
    if removed.empty:
        raise ValueError(f"No existe el experimento {requested} en el registro activo.")

    official_name = text(removed.iloc[0]["Experimento"])
    remaining = current[~mask].copy().reset_index(drop=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    archive_dir = project / "experimentos_eliminados" / f"{safe_name(official_name)}_{stamp}"
    archive_dir.mkdir(parents=True, exist_ok=False)

    moved = []
    sources = (
        (project / "experimentos" / safe_name(official_name), archive_dir / "experimento"),
        (project / "control_roi" / safe_name(official_name), archive_dir / "control_roi"),
    )
    try:
        for source, destination in sources:
            if source.exists():
                shutil.move(str(source), str(destination))
                moved.append((destination, source))

        write_registry(project, remaining)
        _sync_experiment_registries(project, remaining)
        counts_path = write_counts_workbook(project, remaining)
        (archive_dir / "ELIMINADO.txt").write_text(
            "EXPERIMENTO RETIRADO DEL PROYECTO ACTIVO\n"
            f"Fecha: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n"
            f"Experimento: {official_name}\n"
            f"Stacks retirados: {len(removed)}\n"
            f"Animales: {removed['Animal'].nunique()}\n\n"
            "Los datos se conservaron aquí para una recuperación manual. "
            "Las corridas históricas no se modificaron.\n",
            encoding="utf-8",
        )
    except Exception:
        for destination, source in reversed(moved):
            if destination.exists():
                source.parent.mkdir(parents=True, exist_ok=True)
                shutil.move(str(destination), str(source))
        try:
            write_registry(project, current)
            _sync_experiment_registries(project, current)
            write_counts_workbook(project, current)
        except Exception:
            pass
        raise

    return {
        "Experimento": official_name,
        "N_stacks": len(removed),
        "N_animales": int(removed["Animal"].nunique()),
        "Archivo": archive_dir,
        "Conteos": counts_path,
        "Data": remaining,
    }


def save_experiment(project, staging, overwrite=False):
    project = Path(project)
    project.mkdir(parents=True, exist_ok=True)
    staging, staging_duplicates = consolidate_duplicate_stacks(project, staging)
    if staging.empty:
        raise ValueError("El pre-registro está vacío.")
    experiments = {text(value) for value in staging["Experimento"] if text(value)}
    if len(experiments) != 1:
        raise ValueError("El pre-registro debe contener exactamente un experimento.")
    experiment = next(iter(experiments))
    current, current_duplicates = consolidate_duplicate_stacks(project, load_registry(project))
    staging_keys = {stack_id_key(value) for value in staging["Stack_ID"] if stack_id_key(value)}
    current_keys = {stack_id_key(value) for value in current["Stack_ID"] if stack_id_key(value)}
    conflict_keys = staging_keys & current_keys
    conflicts = sorted(
        {
            text(value) for value in staging["Stack_ID"]
            if stack_id_key(value) in conflict_keys
        }
    )
    if conflicts and not overwrite:
        raise DuplicateStackError(conflicts)

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    saved_rows = []
    for _, source_row in staging.iterrows():
        record = refresh_qc(apply_mounting_correction(source_row.to_dict()))
        stack_id = text(record["Stack_ID"])
        destination = project / "experimentos" / safe_name(experiment) / "entradas" / safe_name(stack_id) / stamp
        destination.mkdir(parents=True, exist_ok=True)
        for key, label in (
            ("Archivo_objetos", "objetos"),
            ("Archivo_area", "area_volumen"),
            ("Archivo_ROI", "roi"),
        ):
            source = resolve_project_file_path(project, record.get(key), record, key)
            if source.is_file():
                target = destination / f"{label}{source.suffix.lower()}"
                shutil.copy2(source, target)
                record[key] = portable_project_file_path(project, target)
            else:
                record[key] = text(record.get(key))
        record["Estado_registro"] = "Guardado"
        record["Actualizado"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        saved_rows.append({column: record.get(column, "") for column in META_COLUMNS})

    saved = pd.DataFrame(saved_rows, columns=META_COLUMNS)
    if conflicts:
        current = current[
            ~current["Stack_ID"].map(stack_id_key).isin(conflict_keys)
        ]
    combined = saved.copy() if current.empty else pd.concat([current, saved], ignore_index=True)
    combined, combined_duplicates = consolidate_duplicate_stacks(project, combined)
    registry_path = write_registry(project, combined)
    experiment_file = _write_experiment_registry(project, experiment, combined)
    counts_path = write_counts_workbook(project, combined)
    refreshed_data = load_registry(project).reset_index(drop=True)
    return {
        "Experimento": experiment,
        "N_stacks": len(saved),
        "Reemplazados": conflicts,
        "Registro": registry_path,
        "Registro_experimento": experiment_file,
        "Conteos": counts_path,
        "Data": refreshed_data,
        "Duplicados_reparados": (
            staging_duplicates["Filas_eliminadas"]
            + current_duplicates["Filas_eliminadas"]
            + combined_duplicates["Filas_eliminadas"]
        ),
    }
