import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import numpy as np
import pandas as pd
from openpyxl import load_workbook

import batch_project as batch
import cfos_core as core


class HeatmapV11Tests(unittest.TestCase):

    def test_calculated_filter_uses_surface_over_volume_not_third_column(self):
        frame = pd.DataFrame({
            "Volume (micron^3)": [34.113, 34.854, 60.0],
            "Surface (micron^2)": [93.588, 113.321, 400.0],
            "Nb of obj. voxels": [46, 47, 999],
            "Filtro": ["exclude", "include", "exclude"],
        })
        result = core.calculated_filter(frame).str.casefold().tolist()
        # 34.113: Surface/Volume ~= 2.743 -> incluye aunque la 3a columna sea >2.8.
        # 34.854: Surface/Volume ~= 3.251 -> excluye.
        # 60: no se excluye porque Volume no es <50.
        self.assertEqual(result, ["incluir", "excluir", "incluir"])

    def test_group_map_scales_and_axis_labels_are_fixed(self):
        self.assertEqual(core.GROUP_2D_VMAX_MM3, 80_000.0)
        self.assertEqual(core.GROUP_3D_VMAX_MM3, 300_000.0)
        self.assertEqual(core.X_AXIS_LABEL, "X normalizada")
        self.assertEqual(core.Y_AXIS_LABEL, "Y normalizada")

    def test_difference_scale_uses_pooled_nonzero_percentile(self):
        first = np.array([[0.0, 10.0], [20.0, np.nan]])
        second = np.array([[0.0, 30.0], [40.0, 50.0]])
        expected = np.percentile([10.0, 20.0, 30.0, 40.0, 50.0], 99.5)
        self.assertAlmostEqual(core.pooled_difference_vmax([first, second]), expected)

    def base_record(self, mounted_right):
        return {
            "Stack_ID": "HB_04_R1_C5_BLAa", "Experimento": "HB_04", "Animal": "R1",
            "Corte": "C5", "Region": "BLAa", "Grupo": "SCOP-HB", "Hemisferio": "L",
            "Bregma_mm": -2.3, "Montado_lado_derecho": mounted_right,
            "Archivo_objetos": "objects.xlsx", "Archivo_area": "area.xlsx", "Archivo_ROI": "roi.roi",
            "QC": "OK", "Usar_mapa_grupal": "Si", "Revision_ROI": "Revisado",
        }

    def test_mounting_rules_are_exclusive(self):
        left = core.apply_mounting_correction(self.base_record("No"))
        right = core.apply_mounting_correction(self.base_record("Si"))
        self.assertEqual((left["Flip_X"], left["Flip_Z"], left["Flip_Y"]), ("Si", "No", "No"))
        self.assertEqual((right["Flip_X"], right["Flip_Z"], right["Flip_Y"]), ("No", "Si", "No"))

    def test_unchanged_mounting_preserves_roi_review(self):
        record = {
            "Montado_lado_derecho": "No", "Revision_ROI": "Revisado",
            "QC": "Excluir",
        }
        updated = batch.set_mounting(record, "No")
        self.assertEqual(updated["Revision_ROI"], "Revisado")

    def test_prepare_stack_keeps_z_in_um_and_reverses_only_right_mount(self):
        cells = pd.DataFrame({
            "X": [2.0, 8.0], "Y": [2.0, 8.0], "Z": [1.0, 9.0],
            "Filtro_aplicado": ["incluir", "incluir"],
        })
        polygon = np.array([[0, 0], [10, 0], [10, 10], [0, 10]], dtype=float)
        mocks = (
            patch.object(core, "load_cells", return_value=cells),
            patch.object(core, "load_area", return_value={"Area_um2": 100.0, "Volumen_um3": 1000.0}),
            patch.object(core, "read_roi_polygons", return_value=[polygon]),
        )
        with mocks[0], mocks[1], mocks[2]:
            left = core.prepare_stack(self.base_record("No"), bins2d=4)
            right = core.prepare_stack(self.base_record("Si"), bins2d=4)
        self.assertNotIn("Z_norm", left["valid"].columns)
        self.assertEqual(left["valid"]["Z_corregida_um"].tolist(), [1.0, 9.0])
        self.assertEqual(right["valid"]["Z_corregida_um"].tolist(), [9.0, 1.0])
        self.assertTrue(np.allclose(left["volume2d"], 62.5))
        self.assertEqual(left["valid"]["X_norm"].round(2).tolist(), [0.8, 0.2])
        self.assertEqual(right["valid"]["X_norm"].round(2).tolist(), [0.2, 0.8])
        # Y se exporta anatómicamente: 0=ventral y 1=dorsal.
        self.assertEqual(left["valid"]["Y_norm"].round(2).tolist(), [0.8, 0.2])
        self.assertEqual(right["valid"]["Y_norm"].round(2).tolist(), [0.8, 0.2])

    def test_sector_volumes_reconcile_to_total_volume(self):
        valid = pd.DataFrame({
            "X_norm": [.2, .2, .2, .8],
            "Y_norm": [.2, .2, .8, .8],
        })
        stack = {
            "record": {"QC": "OK", "Experimento": "HB_04", "Animal": "R1", "Grupo": "SCOP-HB", "Region": "BLAa", "Volumen_um3": 800.0},
            "valid": valid,
            "roi_fraction": np.ones((4, 4)),
        }
        result = core.sector_summaries([stack])
        self.assertAlmostEqual(result["Volumen_um3"].sum(), 800.0)
        self.assertEqual(result["Celulas_cFos"].sum(), 4)
        by_sector = result.set_index("Sector")["Celulas_cFos"].to_dict()
        self.assertEqual(by_sector["Medial-Ventral"], 2)
        self.assertEqual(by_sector["Lateral-Ventral"], 0)
        self.assertEqual(by_sector["Medial-Dorsal"], 1)
        self.assertEqual(by_sector["Lateral-Dorsal"], 1)

    def test_excel_spatial_headers_state_anatomical_convention(self):
        frame = pd.DataFrame({
            "X_norm": [.25],
            "Y_norm": [.75],
            "Centroide_Y_norm": [.7],
            "Hotspot_3D_Y_norm": [.8],
        })
        exported = core.spatial_export_frame(frame)
        self.assertIn("X_ML_norm_0_medial_1_lateral", exported.columns)
        self.assertIn("Y_DV_norm_0_ventral_1_dorsal", exported.columns)
        self.assertIn("Centroide_Y_DV_norm_0_ventral_1_dorsal", exported.columns)
        self.assertIn("Hotspot_3D_Y_DV_norm_0_ventral_1_dorsal", exported.columns)
        self.assertAlmostEqual(exported.iloc[0]["Y_DV_norm_0_ventral_1_dorsal"], .75)

    def test_final_workbook_exports_corrected_y_and_coordinate_dictionary(self):
        cells = pd.DataFrame({
            "X": [2.0, 8.0], "Y": [2.0, 8.0], "Z": [1.0, 9.0],
            "Filtro_aplicado": ["incluir", "incluir"],
        })
        polygon = np.array([[0, 0], [10, 0], [10, 10], [0, 10]], dtype=float)
        with patch.object(core, "load_cells", return_value=cells), patch.object(
            core, "load_area", return_value={"Area_um2": 100.0, "Volumen_um3": 1000.0}
        ), patch.object(core, "read_roi_polygons", return_value=[polygon]):
            stack = core.prepare_stack(self.base_record("No"), bins2d=4)

        with tempfile.TemporaryDirectory() as directory:
            project = Path(directory)
            registry = pd.DataFrame([core.apply_mounting_correction(self.base_record("No"))])
            registry = registry.reindex(columns=core.META_COLUMNS, fill_value="")
            registry.to_excel(project / "registro_general_stacks.xlsx", index=False)
            with patch.object(core, "prepare_stack", return_value=stack):
                run_dir = core.run_final(project)
            workbook = load_workbook(run_dir / "resultados_cfos_2D_3D.xlsx", read_only=True, data_only=True)
            try:
                self.assertIn("Convenciones_espaciales", workbook.sheetnames)
                sheet = workbook["Celulas_3D"]
                headers = [cell.value for cell in sheet[1]]
                y_header = "Y_DV_norm_0_ventral_1_dorsal"
                self.assertIn(y_header, headers)
                y_column = headers.index(y_header) + 1
                values = [sheet.cell(row, y_column).value for row in range(2, sheet.max_row + 1)]
                self.assertEqual([round(value, 2) for value in values], [0.8, 0.2])
                version_headers = [cell.value for cell in workbook["Version"][1]]
                self.assertIn("Convencion_Y", version_headers)
                self.assertIn("Escala_2D_grupal_vmax_celulas_mm3", version_headers)
                self.assertIn("Escala_3D_grupal_vmax_celulas_mm3", version_headers)
            finally:
                workbook.close()

    def test_bregma_plot_outputs_three_group_design(self):
        rows = []
        for group, base in [("Vehículo", 8), ("SCOP-HB", 11), ("LiCl", 16)]:
            for animal in ("R1", "R2", "R3"):
                for bregma in (-2.7, -2.3, -1.9):
                    rows.append({"Experimento": "X", "Animal": animal, "Grupo": group, "Region": "BLAa", "Bregma_mm": bregma, "Densidad_x10_3_mm3": base, "Densidad_celulas_mm3": base * 1000})
        bregma = pd.DataFrame(rows)
        animal = bregma.groupby(["Experimento", "Animal", "Grupo", "Region"], as_index=False).first()
        with tempfile.TemporaryDirectory() as directory:
            files = core.plot_density_figures(animal, bregma, directory)
            expected = Path(directory) / "graficas_suplementarias" / "Densidad_por_bregma_3_barras_BLAa.png"
            self.assertTrue(expected.is_file())
            self.assertGreater(expected.stat().st_size, 10_000)
            self.assertIn(str(expected), files)

    def test_prism_sheet_has_one_column_per_group_and_region(self):
        rows = []
        for group in ("Vehículo", "SCOP-HB", "LiCl"):
            for region in ("BLAa", "BLAp"):
                rows.append({
                    "Experimento": "X", "Animal": "R1", "Grupo": group,
                    "Region": region, "Densidad_x10_3_mm3": 12.0,
                })
        result = core.prism_wide_density(pd.DataFrame(rows))
        self.assertEqual(result.shape, (1, 6))
        self.assertIn("Vehículo_BLAa", result.columns)
        self.assertIn("LiCl_BLAp", result.columns)


if __name__ == "__main__":
    unittest.main(verbosity=2)
