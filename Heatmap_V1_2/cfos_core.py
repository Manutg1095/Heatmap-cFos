"""Motor reproducible para análisis espacial c-Fos en BLA.

Principios:
- nunca modifica los archivos fuente;
- la unidad inferencial es el animal;
- los conteos y volúmenes se suman antes de calcular densidad;
- BLAa y BLAp usan espacios AP independientes;
- lo no muestreado permanece como NaN, no como densidad cero.
"""
from __future__ import annotations

import re
import unicodedata
from collections import defaultdict
from datetime import datetime
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.path import Path as MplPath
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

APP_VERSION = "1.2.4"
RANGES = {"BLAa": (-1.8, -3.3), "BLAp": (-3.6, -4.8)}
X_COORDINATE_CONVENTION = "0=medial; 1=lateral"
Y_COORDINATE_CONVENTION = "0=ventral; 1=dorsal"
Y_COORDINATE_TRANSFORM = "Y_DV_norm = 1 - Y_imagen_norm"
X_AXIS_LABEL = "X normalizada"
Y_AXIS_LABEL = "Y normalizada"
SPATIAL_EXPORT_RENAMES = {
    "X_norm": "X_ML_norm_0_medial_1_lateral",
    "Y_norm": "Y_DV_norm_0_ventral_1_dorsal",
    "Centroide_X_norm": "Centroide_X_ML_norm_0_medial_1_lateral",
    "Centroide_Y_norm": "Centroide_Y_DV_norm_0_ventral_1_dorsal",
    "SD_X_norm": "SD_X_ML_norm",
    "SD_Y_norm": "SD_Y_DV_norm",
    "Hotspot_X_norm": "Hotspot_X_ML_norm_0_medial_1_lateral",
    "Hotspot_Y_norm": "Hotspot_Y_DV_norm_0_ventral_1_dorsal",
    "Hotspot_2D_X_norm": "Hotspot_2D_X_ML_norm_0_medial_1_lateral",
    "Hotspot_2D_Y_norm": "Hotspot_2D_Y_DV_norm_0_ventral_1_dorsal",
    "Hotspot_3D_X_norm": "Hotspot_3D_X_ML_norm_0_medial_1_lateral",
    "Hotspot_3D_Y_norm": "Hotspot_3D_Y_DV_norm_0_ventral_1_dorsal",
}
META_COLUMNS = [
    "Stack_ID", "Experimento", "Animal", "Corte", "Region", "Grupo",
    "Fuente_grupo", "Hemisferio", "Bregma_mm", "Fuente_bregma",
    "Bregma_area_mm", "Montado_lado_derecho", "Direccion_Z_AP",
    "Flip_X", "Flip_Y", "Flip_Z", "Revision_ROI",
    "QC", "Usar_mapa_grupal", "Incluir_analisis", "Archivo_objetos", "Archivo_area",
    "Archivo_ROI", "Area_um2", "Volumen_um3", "Objetos_XYZ",
    "Objetos_incluidos_filtro", "Objetos_fuera_ROI", "Objetos_validos",
    "Alineacion_XY", "Estado_registro", "Detalle_QC_base", "Detalle_QC",
    "Notas", "Actualizado",
]

BINS_2D = 30
BINS_XY_3D = 20
AP_VOXEL_MM = 0.025
AP_MARGIN_MM = 0.15
LOCAL_Z_BINS = 12
GROUP_2D_VMAX_MM3 = 80_000.0
GROUP_3D_VMAX_MM3 = 300_000.0
GROUP_COLORS = {"Vehículo": "#C9C6C6", "SCOP-HB": "#4C95C7", "LiCl": "#E53935"}


def text(value):
    return "" if pd.isna(value) else str(value).strip()


def truthy(value):
    return text(value).lower() in {"si", "sí", "yes", "true", "1"}


def selected_for_analysis(value):
    """Los registros antiguos sin esta columna se consideran seleccionados."""
    return text(value) == "" or truthy(value)


def yes_no(value):
    normalized = normalized_label(value)
    if normalized in {"si", "yes", "true", "1", "derecho", "right"}:
        return "Si"
    if normalized in {"no", "false", "0", "izquierdo", "left"}:
        return "No"
    return ""


def apply_mounting_correction(record):
    """Deriva las correcciones espaciales desde una sola decisión auditable.

    Todos los tejidos son hemisferio izquierdo. Si el montaje fue izquierdo se
    corrige X; si fue derecho se conserva X y se invierte el orden óptico Z.
    Después de corregir, Z siempre aumenta de rostral a caudal (Posterior).
    """
    result = dict(record)
    mounting = yes_no(result.get("Montado_lado_derecho"))
    result["Hemisferio"] = "L"
    result["Flip_Y"] = "No"
    result["Direccion_Z_AP"] = "Posterior"
    if mounting == "Si":
        result["Montado_lado_derecho"] = "Si"
        result["Flip_X"] = "No"
        result["Flip_Z"] = "Si"
    elif mounting == "No":
        result["Montado_lado_derecho"] = "No"
        result["Flip_X"] = "Si"
        result["Flip_Z"] = "No"
    else:
        result["Montado_lado_derecho"] = ""
        result["Flip_X"] = "No"
        result["Flip_Z"] = "No"
    return result


def safe_name(value):
    return re.sub(r"[^A-Za-z0-9._-]+", "_", text(value)).strip("_") or "sin_nombre"


def portable_project_file_path(project, value):
    """Convierte una copia interna en una ruta relativa y transportable."""
    raw = text(value)
    if not raw:
        return ""
    project = Path(project).expanduser().resolve()
    path = Path(raw).expanduser()
    try:
        resolved = path.resolve()
        return resolved.relative_to(project).as_posix()
    except (OSError, ValueError):
        return raw


def resolve_project_file_path(project, value, record=None, key=None):
    """Localiza una entrada aunque el proyecto se haya movido de carpeta."""
    raw = text(value)
    if not raw:
        return Path("")
    project = Path(project).expanduser().resolve()
    original = Path(raw).expanduser()

    candidates = []
    # Si el registro proviene de otra instalación, la copia perteneciente al
    # proyecto activo tiene prioridad incluso cuando la carpeta antigua todavía
    # existe. Así una corrida nunca vuelve a leer silenciosamente la app pasada.
    parts = original.parts
    if "experimentos" in parts:
        marker = parts.index("experimentos")
        candidates.append(project.joinpath(*parts[marker:]))
    if original.is_absolute():
        candidates.append(original)
    else:
        candidates.append(project / original)

    for candidate in candidates:
        if candidate.is_file():
            return candidate.resolve()

    # Último recurso para registros antiguos: usa la copia interna más reciente
    # del mismo stack y del mismo tipo de archivo.
    record = record or {}
    experiment = safe_name(record.get("Experimento", ""))
    stack_id = safe_name(record.get("Stack_ID", ""))
    labels = {
        "Archivo_objetos": ("objetos", {".xlsx", ".xls"}),
        "Archivo_area": ("area_volumen", {".xlsx", ".xls"}),
        "Archivo_ROI": ("roi", {".roi", ".zip"}),
    }
    if key in labels and experiment and stack_id:
        label, suffixes = labels[key]
        stack_root = project / "experimentos" / experiment / "entradas" / stack_id
        internal = sorted(
            (
                path for path in stack_root.glob(f"*/{label}.*")
                if path.is_file() and path.suffix.lower() in suffixes
            ),
            key=lambda path: path.parent.name,
            reverse=True,
        )
        if internal:
            return internal[0].resolve()
    return original


def resolve_record_file_paths(project, record):
    """Devuelve un registro con sus tres archivos resueltos en el proyecto."""
    output = dict(record)
    for key in ("Archivo_objetos", "Archivo_area", "Archivo_ROI"):
        resolved = resolve_project_file_path(project, output.get(key), output, key)
        output[key] = str(resolved) if text(output.get(key)) else ""
    return output


def save_figure_atomic(figure, filename, **kwargs):
    """Guarda una figura sin exponer un PNG/PDF parcial como resultado final."""
    path = Path(filename)
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.name}.partial")
    image_format = path.suffix.lstrip(".").lower()
    last_error = None
    try:
        for _ in range(2):
            try:
                if temporary.exists():
                    temporary.unlink()
                figure.savefig(temporary, format=image_format, **kwargs)
                if image_format == "png":
                    # La lectura fuerza la validación de los chunks/CRC antes
                    # de publicar el archivo. Si falla, se reintenta una vez.
                    plt.imread(temporary)
                temporary.replace(path)
                return path
            except Exception as exc:
                last_error = exc
        raise RuntimeError(f"No se pudo guardar una figura íntegra: {path}") from last_error
    finally:
        if temporary.exists():
            temporary.unlink()


def add_fixed_density_colorbar(figure, mappable, axis, vmax, **kwargs):
    """Añade una barra fija y marca explícitamente la saturación superior."""
    colorbar = figure.colorbar(
        mappable,
        ax=axis,
        extend="max",
        label="c-Fos+ cells/mm³",
        **kwargs,
    )
    ticks = np.linspace(0, float(vmax), 5)
    labels = [f"{value:,.0f}" for value in ticks]
    labels[-1] = f"≥{float(vmax):,.0f}"
    colorbar.set_ticks(ticks)
    colorbar.set_ticklabels(labels)
    return colorbar


def findcol(df, *names):
    lookup = {str(column).strip().lower(): column for column in df.columns}
    return next((lookup[name.lower()] for name in names if name.lower() in lookup), None)


def normalized_label(value):
    value = unicodedata.normalize("NFKD", text(value)).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", "", value)


def numeric_value(value):
    if pd.isna(value):
        return None
    if isinstance(value, (int, float, np.integer, np.floating)):
        return float(value) if np.isfinite(value) else None
    value = text(value).replace("−", "-").replace(",", ".")
    match = re.search(r"[-+]?\d+(?:\.\d+)?", value)
    return float(match.group()) if match else None


def read_table(file):
    """Lee un catálogo CSV/Excel sin modificar el archivo fuente."""
    path = Path(file)
    if path.suffix.lower() == ".csv":
        return pd.read_csv(path, dtype=str).fillna("")
    if path.suffix.lower() in {".xlsx", ".xls"}:
        return pd.read_excel(path, dtype=str).fillna("")
    raise ValueError("El catálogo debe ser CSV, XLSX o XLS.")


def normalized_experiment(value):
    value = text(value).replace("-", "_").replace(" ", "_")
    value = re.sub(r"_+", "_", value).strip("_")
    hb = re.fullmatch(r"(?i)HB_?(\d+)", value)
    if hb:
        return f"HB_{int(hb.group(1)):02d}"
    scop = re.fullmatch(r"(?i)SCOP_?LiCl_?(\d+)", value)
    if scop:
        return f"SCOP_LiCl_{int(scop.group(1)):02d}"
    return value


def normalized_animal(value):
    match = re.search(r"(?i)R\s*0*(\d+)", text(value))
    return f"R{int(match.group(1))}" if match else ""


def normalized_cut(value):
    match = re.search(r"(?i)C\s*0*(\d+)", text(value))
    return f"C{int(match.group(1))}" if match else ""


def normalized_region(value):
    match = re.search(r"(?i)(BLAa|BLAp)", text(value))
    if not match:
        return ""
    return "BLAa" if match.group(1).lower() == "blaa" else "BLAp"


def load_group_catalog(file):
    frame = read_table(file)
    columns = {normalized_label(column): column for column in frame.columns}
    required = {"experimento", "animal", "grupo"}
    if not required.issubset(columns):
        raise ValueError("La asignación debe tener las columnas Experimento, Animal y Grupo.")
    output = pd.DataFrame(
        {
            "Experimento": frame[columns["experimento"]].map(normalized_experiment),
            "Animal": frame[columns["animal"]].map(normalized_animal),
            "Grupo": frame[columns["grupo"]].map(text),
        }
    )
    incomplete = output.eq("").any(axis=1)
    if incomplete.any():
        raise ValueError(f"La asignación contiene {int(incomplete.sum())} fila(s) incompleta(s).")
    duplicated = output.duplicated(["Experimento", "Animal"], keep=False)
    if duplicated.any():
        keys = output.loc[duplicated, ["Experimento", "Animal"]].drop_duplicates()
        detail = ", ".join(keys.astype(str).agg("/".join, axis=1))
        raise ValueError(f"Hay asignaciones duplicadas para: {detail}.")
    return output.sort_values(["Grupo", "Experimento", "Animal"]).reset_index(drop=True)


def load_bregma_catalog(file):
    frame = read_table(file)
    columns = {normalized_label(column): column for column in frame.columns}
    experiment = columns.get("experimento")
    animal_cut = columns.get("animalycorte")
    region = columns.get("region")
    bregma = columns.get("bregma")
    if not all((experiment, animal_cut, region, bregma)):
        raise ValueError("El registro debe tener Experimento, Animal y corte, Región y Bregma.")
    output = pd.DataFrame(
        {
            "Experimento": frame[experiment].map(normalized_experiment),
            "Animal": frame[animal_cut].map(normalized_animal),
            "Corte": frame[animal_cut].map(normalized_cut),
            "Region": frame[region].map(normalized_region),
            "Bregma_mm": frame[bregma].map(numeric_value),
        }
    )
    missing = output[["Experimento", "Animal", "Corte", "Region"]].eq("").any(axis=1)
    missing |= output["Bregma_mm"].isna()
    if missing.any():
        raise ValueError(
            f"El registro de bregma contiene {int(missing.sum())} fila(s) incompleta(s) o inválida(s)."
        )
    duplicated = output.duplicated(["Experimento", "Animal", "Corte", "Region"], keep=False)
    if duplicated.any():
        keys = output.loc[duplicated, ["Experimento", "Animal", "Corte", "Region"]].drop_duplicates()
        detail = ", ".join(keys.astype(str).agg("/".join, axis=1))
        raise ValueError(f"Hay bregmas duplicados para: {detail}.")
    outside = []
    for row in output.itertuples():
        anterior, posterior = RANGES[row.Region]
        if not posterior <= float(row.Bregma_mm) <= anterior:
            outside.append(f"{row.Experimento}/{row.Animal}/{row.Corte}/{row.Region}={row.Bregma_mm:g}")
    if outside:
        raise ValueError("Hay bregmas fuera del rango anatómico configurado: " + ", ".join(outside))
    return output.sort_values(["Experimento", "Animal", "Region", "Corte"]).reset_index(drop=True)


def load_bregma(file, required=True):
    """Lee un bregma único desde cualquier hoja del Excel de área/volumen."""
    sheets = pd.read_excel(file, sheet_name=None)
    found = []
    for sheet_name, frame in sheets.items():
        for column in frame.columns:
            if "bregma" not in normalized_label(column):
                continue
            for row, value in frame[column].items():
                number = numeric_value(value)
                if number is not None:
                    found.append((number, sheet_name, str(column), int(row) + 2))

    if not found:
        raw_sheets = pd.read_excel(file, sheet_name=None, header=None)
        for sheet_name, frame in raw_sheets.items():
            values = frame.to_numpy(dtype=object)
            for row in range(values.shape[0]):
                for col in range(values.shape[1]):
                    if "bregma" not in normalized_label(values[row, col]):
                        continue
                    neighbours = []
                    if col + 1 < values.shape[1]:
                        neighbours.append(values[row, col + 1])
                    if row + 1 < values.shape[0]:
                        neighbours.append(values[row + 1, col])
                    for value in neighbours:
                        number = numeric_value(value)
                        if number is not None:
                            found.append((number, sheet_name, f"R{row+1}C{col+1}", row + 1))
                            break

    unique = {round(item[0], 6) for item in found}
    if not unique:
        if required:
            raise ValueError("No encontré un valor de bregma en el Excel de área/volumen.")
        return None
    if len(unique) > 1:
        values = ", ".join(f"{value:g}" for value in sorted(unique, reverse=True))
        raise ValueError(f"El Excel de área/volumen contiene más de un bregma distinto: {values}.")
    return found[0][0]


def parse_name(source):
    name = Path(source).name
    patterns = [
        (r"(?i)(HB)[_ -]?(\d+)[_ -]+R(\d+)[_ -]+C(\d+)[_ -]+(BLAa|BLAp)", "HB"),
        (r"(?i)(SCOP[_ -]?LiCl)[_ -]?(\d+)[_ -]+R(\d+)[_ -]+C(\d+)[_ -]+(BLAa|BLAp)", "SCOP_LiCl"),
    ]
    match = prefix = None
    for pattern, canonical in patterns:
        match = re.search(pattern, name)
        if match:
            prefix = canonical
            break
    if not match:
        raise ValueError(
            "El nombre debe contener HB_XX_RXX_CX_BLAa/BLAp o "
            "SCOP_LiCl_XX_RXX_CX_BLAa/BLAp."
        )
    _, number, animal, cut, region = match.groups()
    experiment = f"{prefix}_{int(number):02d}"
    animal = f"R{int(animal)}"
    cut = f"C{int(cut)}"
    region = normalized_region(region)
    return {
        "Experimento": experiment,
        "Animal": animal,
        "Corte": cut,
        "Region": region,
        "Stack_ID": f"{experiment}_{animal}_{cut}_{region}",
    }


def calculated_filter(df):
    """Aplica el filtro morfológico directamente desde el Excel de objetos.

    Se excluye un objeto únicamente cuando se cumplen simultáneamente:
    - Volume (micron^3) < 50
    - Surface (micron^2) / Volume (micron^3) > 2.8

    No se usan posiciones de columna (A/C) ni una columna Filtro preexistente,
    para que el resultado sea reproducible aunque cambie el orden de columnas.
    """
    volume_column = findcol(
        df,
        "Volume (micron^3)", "Volume", "Volume um3",
        "Volumen", "Volumen um3", "Volumen (micron^3)", "Volumen (µm3)",
    )
    surface_column = findcol(
        df,
        "Surface (micron^2)", "Surface", "Surface um2",
        "Superficie", "Superficie um2", "Superficie (micron^2)", "Superficie (µm2)",
    )
    if volume_column is None or surface_column is None:
        raise ValueError(
            "El Excel de objetos requiere columnas de Volume (micron^3) y "
            "Surface (micron^2) para aplicar el filtro morfológico."
        )

    volume = pd.to_numeric(df[volume_column], errors="coerce")
    surface = pd.to_numeric(df[surface_column], errors="coerce")
    surface_volume = surface / volume

    exclude = (volume < 50) & (surface_volume > 2.8)
    return pd.Series(np.where(exclude, "Excluir", "incluir"), index=df.index)


def load_cells(file):
    raw = pd.read_excel(file)
    columns = [findcol(raw, axis) for axis in ("X", "Y", "Z")]
    if any(column is None for column in columns):
        raise ValueError("El Excel de objetos requiere las columnas X, Y y Z.")
    cells = raw[columns].copy()
    cells.columns = ["X", "Y", "Z"]
    cells["Filtro_aplicado"] = calculated_filter(raw)
    for column in ("X", "Y", "Z"):
        cells[column] = pd.to_numeric(cells[column], errors="coerce")
    return cells.dropna(subset=["X", "Y", "Z"]).reset_index(drop=True)


def load_area(file):
    if not file:
        return {}
    sheets = pd.read_excel(file, sheet_name=None)
    output = {}
    bregma = load_bregma(file, required=False)
    if bregma is not None:
        output["Bregma_area_mm"] = bregma
    requested = {
        "Area_um2": ("Area", "Area (micron^2)", "Area um2", "Área", "Área um2"),
        "Volumen_um3": (
            "Volume", "Volume (micron^3)", "Volume um3", "Volumen",
            "Volumen um3", "Volumen (micron^3)", "Volumen (µm3)",
        ),
    }
    for frame in sheets.values():
        for key, names in requested.items():
            if key in output:
                continue
            column = findcol(frame, *names)
            if column is not None and len(frame):
                values = pd.to_numeric(frame[column], errors="coerce").dropna()
                if len(values):
                    output[key] = float(values.iloc[0])
    return output


def polygon_area(points):
    points = np.asarray(points, dtype=float)
    if len(points) < 3:
        return 0.0
    x, y = points[:, 0], points[:, 1]
    return abs(np.dot(x, np.roll(y, 1)) - np.dot(y, np.roll(x, 1))) / 2


def read_roi_polygons(file):
    """Devuelve contornos ImageJ como arreglos Nx2 en coordenadas de la ROI."""
    try:
        from roifile import roiread
    except ImportError as exc:
        raise RuntimeError("Falta roifile. Instala las dependencias con: python -m pip install -r requirements.txt") from exc

    loaded = roiread(file)
    rois = loaded if isinstance(loaded, (list, tuple)) else [loaded]
    polygons = []
    for roi in rois:
        try:
            coordinates = np.asarray(roi.coordinates(), dtype=float)
        except Exception:
            coordinates = np.empty((0, 2), dtype=float)
        candidates = [coordinates] if coordinates.ndim == 2 else list(coordinates) if coordinates.ndim == 3 else []
        for candidate in candidates:
            candidate = np.asarray(candidate, dtype=float)
            if candidate.ndim == 2 and candidate.shape[1] >= 2 and len(candidate) >= 3:
                polygons.append(candidate[:, :2])
        if not candidates or not any(len(candidate) >= 3 for candidate in candidates):
            left, top = float(roi.left), float(roi.top)
            right, bottom = float(roi.right), float(roi.bottom)
            if right > left and bottom > top:
                polygons.append(np.array([[left, top], [right, top], [right, bottom], [left, bottom]], dtype=float))
    if not polygons:
        raise ValueError("La ROI no contiene un polígono/rectángulo utilizable.")
    return polygons


def points_inside(points, polygons):
    points = np.asarray(points, dtype=float)
    inside = np.zeros(len(points), dtype=bool)
    for polygon in polygons:
        # Path(..., closed=True) trata el último vértice como CLOSEPOLY y puede
        # eliminar una esquina. Cerramos explícitamente el contorno.
        closed = np.vstack([polygon, polygon[0]])
        inside |= MplPath(closed).contains_points(points, radius=1e-9)
    return inside


def orient_xy(points, hemisphere, flip_x=False, flip_y=False):
    points = np.asarray(points, dtype=float).copy()
    # El hemisferio derecho se refleja; Flip_X permite alternar/corregir esa decisión.
    if (text(hemisphere).upper() == "R") ^ bool(flip_x):
        points[:, 0] = 1.0 - points[:, 0]
    if flip_y:
        points[:, 1] = 1.0 - points[:, 1]
    return points


def to_anatomical_y(points):
    """Convierte Y de imagen a Y dorsoventral: 0=ventral, 1=dorsal."""
    points = np.asarray(points, dtype=float).copy()
    points[:, 1] = 1.0 - points[:, 1]
    return points


def spatial_export_frame(frame):
    """Da nombres anatómicos explícitos a las coordenadas exportadas a Excel."""
    return frame.rename(columns={
        source: target
        for source, target in SPATIAL_EXPORT_RENAMES.items()
        if source in frame.columns
    })


def align_cells_and_roi(cells, polygons, hemisphere, flip_x=False, flip_y=False):
    """Expresa células y ROI en coordenadas anatómicas normalizadas 0--1.

    Se aceptan coordenadas con un origen común, o una ROI exportada con un
    desplazamiento de lienzo constante que debe llevarse a (0, 0). Nunca se
    reescala cada conjunto de células a su propio bounding box. X queda como
    0=medial/1=lateral y Y como 0=ventral/1=dorsal.
    """
    all_roi = np.vstack(polygons)
    roi_min, roi_max = all_roi.min(axis=0), all_roi.max(axis=0)
    roi_span = np.where(roi_max > roi_min, roi_max - roi_min, 1.0)
    roi_norm = [(polygon - roi_min) / roi_span for polygon in polygons]

    raw_xy = cells[["X", "Y"]].to_numpy(float)
    direct = (raw_xy - roi_min) / roi_span
    roi_origin = raw_xy / roi_span
    direct_inside = points_inside(direct, roi_norm).mean() if len(direct) else 0.0
    origin_inside = points_inside(roi_origin, roi_norm).mean() if len(roi_origin) else 0.0
    if direct_inside >= origin_inside:
        cell_norm, mode, chosen_fraction = direct, "ROI_e_Imaris_mismo_origen", direct_inside
    else:
        cell_norm, mode, chosen_fraction = roi_origin, "ROI_desplazada_a_origen", origin_inside
    if chosen_fraction < 0.25:
        raise ValueError(
            "La ROI y las coordenadas XY no se pueden registrar sin deformación "
            f"(mejor coincidencia: {chosen_fraction:.1%}). Revisa orientación, unidades y origen."
        )

    inside = points_inside(cell_norm, roi_norm)
    cell_norm = orient_xy(cell_norm, hemisphere, flip_x, flip_y)
    roi_norm = [orient_xy(polygon, hemisphere, flip_x, flip_y) for polygon in roi_norm]
    # Imaris/ImageJ usan Y de imagen (0 arriba). La salida usa una coordenada
    # anatómica cartesiana: 0 ventral y 1 dorsal.
    cell_norm = to_anatomical_y(cell_norm)
    roi_norm = [to_anatomical_y(polygon) for polygon in roi_norm]
    output = cells.copy()
    output["X_norm"] = cell_norm[:, 0]
    output["Y_norm"] = cell_norm[:, 1]
    output["Dentro_ROI"] = inside
    output["Filtro_incluye"] = output["Filtro_aplicado"].astype(str).str.lower().isin(["include", "incluir"])
    return output, roi_norm, mode, direct_inside, origin_inside


def roi_fraction_grid(polygons, bins=BINS_2D, samples=8):
    fractions = np.zeros((bins, bins), dtype=float)  # y, x
    offsets = (np.arange(samples) + 0.5) / samples
    for yi in range(bins):
        ys = (yi + offsets) / bins
        for xi in range(bins):
            xs = (xi + offsets) / bins
            xx, yy = np.meshgrid(xs, ys)
            fractions[yi, xi] = points_inside(np.column_stack([xx.ravel(), yy.ravel()]), polygons).mean()
    return fractions


def density_from_counts(counts, volume_um3):
    return np.divide(
        counts * 1e9,
        volume_um3,
        out=np.full(np.shape(counts), np.nan, dtype=float),
        where=np.asarray(volume_um3) > 0,
    )


def pooled_difference_vmax(difference_arrays, percentile=99.5):
    """Límite simétrico común para mapas diferenciales de una modalidad."""
    finite = [
        np.abs(array[np.isfinite(array)])
        for array in difference_arrays
        if np.isfinite(array).any()
    ]
    if not finite:
        return 1.0
    pooled = np.concatenate(finite)
    nonzero = pooled[pooled > 0]
    values = nonzero if len(nonzero) else pooled
    return max(float(np.percentile(values, percentile)), np.finfo(float).eps)


def normalized_entropy(counts):
    counts = np.asarray(counts, dtype=float)
    positive = counts[counts > 0]
    if len(positive) <= 1:
        return 0.0 if len(positive) == 1 else np.nan
    probabilities = positive / positive.sum()
    return float(-(probabilities * np.log(probabilities)).sum() / np.log(len(positive)))


def peak_metrics(density):
    if not np.isfinite(density).any():
        return np.nan, None
    index = np.unravel_index(np.nanargmax(density), density.shape)
    return float(density[index]), index


def effective_bregma(record, area_info=None):
    """Devuelve el bregma efectivo; el registro editable tiene prioridad."""
    area_info = area_info or {}
    bregma_metadata = numeric_value(record.get("Bregma_mm"))
    bregma_area = numeric_value(area_info.get("Bregma_area_mm"))
    if bregma_metadata is None and bregma_area is None:
        raise ValueError(
            "No hay bregma ni en el registro editable ni en el Excel de área."
        )
    return float(bregma_metadata if bregma_metadata is not None else bregma_area)


def prepare_stack(record, bins2d=BINS_2D):
    record = apply_mounting_correction(record)
    if yes_no(record.get("Montado_lado_derecho")) == "":
        raise ValueError("Falta clasificar si el corte fue montado de lado derecho.")
    cells = load_cells(record["Archivo_objetos"])
    area_info = load_area(record["Archivo_area"])
    area_um2 = float(area_info.get("Area_um2", np.nan))
    volume_um3 = float(area_info.get("Volumen_um3", np.nan))
    bregma = effective_bregma(record, area_info)
    if not np.isfinite(area_um2) or area_um2 <= 0:
        raise ValueError("El Excel de área/volumen no contiene un área positiva utilizable.")
    if not np.isfinite(volume_um3) or volume_um3 <= 0:
        raise ValueError("El Excel de área/volumen no contiene un volumen positivo utilizable.")
    region = text(record["Region"])
    if region not in RANGES:
        raise ValueError(f"Región no reconocida: {region}.")

    polygons = read_roi_polygons(record["Archivo_ROI"])
    aligned, polygons_norm, alignment, direct_fraction, relative_fraction = align_cells_and_roi(
        cells,
        polygons,
        record["Hemisferio"],
        truthy(record.get("Flip_X")),
        truthy(record.get("Flip_Y")),
    )
    valid = aligned[aligned["Filtro_incluye"] & aligned["Dentro_ROI"]].copy()
    if valid.empty:
        raise ValueError("Ningún objeto incluido quedó dentro de la ROI después de la alineación.")

    thickness_um = volume_um3 / area_um2
    z_min, z_max = float(cells["Z"].min()), float(cells["Z"].max())
    z_mid_um = (z_min + z_max) / 2
    valid["Z_original_um"] = valid["Z"].astype(float)
    valid["Z_corregida_um"] = (
        z_min + z_max - valid["Z_original_um"]
        if truthy(record.get("Flip_Z")) else valid["Z_original_um"]
    )
    # Tras la corrección, Z aumenta de rostral a caudal. Se conserva en µm.
    valid["AP_mm"] = bregma - (valid["Z_corregida_um"] - z_mid_um) / 1000

    roi_fraction = roi_fraction_grid(polygons_norm, bins2d)
    roi_weight = roi_fraction.sum()
    if roi_weight <= 0:
        raise ValueError("La ROI no cubre ningún bin de la cuadrícula 2D.")
    # La ROI completa se transforma a una cuadrícula rectangular 0–1. Así el
    # heatmap 2D muestra posición relativa y no la silueta de BLAa/BLAp.
    local_volume2d = np.full((bins2d, bins2d), volume_um3 / (bins2d * bins2d))
    counts2d, _, _ = np.histogram2d(
        valid["X_norm"], valid["Y_norm"], bins=bins2d, range=((0, 1), (0, 1))
    )
    counts2d = counts2d.T  # y, x
    density2d = density_from_counts(counts2d, local_volume2d)

    z_edges = np.linspace(z_min, z_max if z_max > z_min else z_min + 1, LOCAL_Z_BINS + 1)
    local_counts3d, _ = np.histogramdd(
        valid[["X_norm", "Y_norm", "Z_corregida_um"]].to_numpy(),
        bins=(bins2d, bins2d, z_edges),
        range=None,
    )
    xy_weights = roi_fraction.T / roi_weight  # x, y
    local_volume3d = volume_um3 * xy_weights[:, :, None] / LOCAL_Z_BINS
    density3d = density_from_counts(local_counts3d, local_volume3d)

    peak2d, peak2d_index = peak_metrics(density2d)
    peak3d, peak3d_index = peak_metrics(density3d)
    warnings = []
    included_by_filter = int(aligned["Filtro_incluye"].sum())
    outside_roi = int((aligned["Filtro_incluye"] & ~aligned["Dentro_ROI"]).sum())
    if included_by_filter and outside_roi / included_by_filter > 0.25:
        warnings.append("Más de 25% de los objetos incluidos quedó fuera de la ROI; revisar alineación XY.")
    anterior, posterior = RANGES[region]
    if not posterior <= bregma <= anterior:
        warnings.append(f"Bregma {bregma:g} fuera del rango configurado para {region}.")

    metadata = {key: text(record.get(key)) for key in META_COLUMNS}
    metadata.update(
        {
            "Bregma_mm": bregma,
            "Area_um2": area_um2,
            "Volumen_um3": volume_um3,
            "Espesor_derivado_um": thickness_um,
            "Z_centro_observado_um": z_mid_um,
            "Objetos_XYZ": len(cells),
            "Objetos_incluidos_filtro": included_by_filter,
            "Objetos_fuera_ROI": outside_roi,
            "Objetos_validos": len(valid),
            "Alineacion_XY": alignment,
            "Fraccion_dentro_directa": direct_fraction,
            "Fraccion_dentro_origen": relative_fraction,
            "Densidad_global_mm3": len(valid) / volume_um3 * 1e9,
            "Pico_2D_mm3": peak2d,
            "Entropia_2D": normalized_entropy(counts2d),
            "Pico_3D_stack_mm3": peak3d,
            "Entropia_3D_stack": normalized_entropy(local_counts3d),
            "Centroide_X_norm": float(valid["X_norm"].mean()),
            "Centroide_Y_norm": float(valid["Y_norm"].mean()),
            "Centroide_AP_mm": float(valid["AP_mm"].mean()),
            "Advertencias": " | ".join(warnings),
        }
    )
    if peak2d_index is not None:
        metadata.update(
            {
                "Hotspot_2D_X_norm": (peak2d_index[1] + 0.5) / bins2d,
                "Hotspot_2D_Y_norm": (peak2d_index[0] + 0.5) / bins2d,
            }
        )
    if peak3d_index is not None:
        metadata.update(
            {
                "Hotspot_3D_X_norm": (peak3d_index[0] + 0.5) / bins2d,
                "Hotspot_3D_Y_norm": (peak3d_index[1] + 0.5) / bins2d,
                "Hotspot_3D_Z_um": float((z_edges[peak3d_index[2]] + z_edges[peak3d_index[2] + 1]) / 2),
            }
        )

    return {
        "record": metadata,
        "all_cells": aligned,
        "valid": valid,
        "polygons": polygons_norm,
        "roi_fraction": roi_fraction,
        "counts2d": counts2d,
        "volume2d": local_volume2d,
        "density2d": density2d,
        "local_counts3d": local_counts3d,
        "local_volume3d": local_volume3d,
        "local_density3d": density3d,
        "z_edges_um": z_edges,
        "thickness_um": thickness_um,
    }


def draw_roi(ax, polygons, color="#202020", linewidth=1.4):
    for polygon in polygons:
        closed = np.vstack([polygon, polygon[0]])
        ax.plot(closed[:, 0], closed[:, 1], color=color, linewidth=linewidth)


def plot_stack_outputs(stack, output_dir):
    output_dir = Path(output_dir)
    for folder in ("QC", "mapas_2D_stack", "mapas_3D_stack"):
        (output_dir / folder).mkdir(parents=True, exist_ok=True)
    sid = safe_name(stack["record"]["Stack_ID"])
    cells, valid = stack["all_cells"], stack["valid"]

    fig, axes = plt.subplots(1, 2, figsize=(11, 5))
    filter_excluded = cells[~cells["Filtro_incluye"]]
    roi_excluded = cells[cells["Filtro_incluye"] & ~cells["Dentro_ROI"]]
    axes[0].scatter(filter_excluded["X_norm"], filter_excluded["Y_norm"], s=8, c="#9a9a9a", label="filtro", alpha=0.55)
    axes[0].scatter(roi_excluded["X_norm"], roi_excluded["Y_norm"], s=9, c="#f3a712", label="fuera ROI", alpha=0.7)
    axes[0].scatter(valid["X_norm"], valid["Y_norm"], s=9, c="#d1495b", label="válidas", alpha=0.75)
    draw_roi(axes[0], stack["polygons"])
    axes[0].set(xlim=(0, 1), ylim=(0, 1), xlabel=X_AXIS_LABEL, ylabel=Y_AXIS_LABEL, title="ROI y objetos")
    axes[0].legend(fontsize=8)
    axes[1].hist(valid["Z_corregida_um"], bins=20, color="#d1495b", alpha=0.75)
    axes[1].set(xlabel="Z corregida (µm)", ylabel="objetos", title="Profundidad óptica corregida")
    fig.suptitle(f"QC — {stack['record']['Stack_ID']} | {stack['record']['Alineacion_XY']}")
    save_figure_atomic(fig, output_dir / "QC" / f"{sid}_QC.png", dpi=180, bbox_inches="tight")
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(6.5, 5.5))
    image = ax.imshow(stack["density2d"], origin="lower", extent=(0, 1, 0, 1), cmap="magma")
    ax.set(
        title=f"2D normalizado — {stack['record']['Stack_ID']}",
        xlabel=X_AXIS_LABEL, ylabel=Y_AXIS_LABEL,
        xlim=(0, 1), ylim=(0, 1), aspect="equal",
    )
    fig.colorbar(image, ax=ax, label="células/mm³")
    save_figure_atomic(fig, output_dir / "mapas_2D_stack" / f"{sid}_2D.png", dpi=180, bbox_inches="tight")
    plt.close(fig)

    density = stack["local_density3d"]
    counts = stack["local_counts3d"]
    indices = np.argwhere((counts > 0) & np.isfinite(density))
    fig = plt.figure(figsize=(7, 6))
    ax = fig.add_subplot(111, projection="3d")
    if len(indices):
        values = density[tuple(indices.T)]
        scatter = ax.scatter(
            (indices[:, 0] + 0.5) / density.shape[0],
            (indices[:, 1] + 0.5) / density.shape[1],
            (stack["z_edges_um"][indices[:, 2]] + stack["z_edges_um"][indices[:, 2] + 1]) / 2,
            c=values,
            s=14 + 4 * counts[tuple(indices.T)],
            cmap="magma",
            alpha=0.75,
        )
        fig.colorbar(scatter, ax=ax, shrink=0.65, label="células/mm³")
    ax.set(xlabel=X_AXIS_LABEL, ylabel=Y_AXIS_LABEL, zlabel="Z corregida (µm)", title=f"3D stack — {stack['record']['Stack_ID']}")
    save_figure_atomic(fig, output_dir / "mapas_3D_stack" / f"{sid}_3D.png", dpi=180, bbox_inches="tight")
    plt.close(fig)


def stack_2d_rows(stack):
    record = stack["record"]
    rows = []
    bins = stack["counts2d"].shape[0]
    for yi, xi in np.argwhere(stack["volume2d"] > 0):
        rows.append(
            {
                "Stack_ID": record["Stack_ID"], "Experimento": record["Experimento"],
                "Animal": record["Animal"], "Grupo": record["Grupo"], "Region": record["Region"],
                "Bregma_mm": record["Bregma_mm"], "Bin_X": int(xi), "Bin_Y": int(yi),
                "X_norm": (xi + 0.5) / bins, "Y_norm": (yi + 0.5) / bins,
                "Celulas": int(stack["counts2d"][yi, xi]),
                "Volumen_um3": float(stack["volume2d"][yi, xi]),
                "Densidad_celulas_mm3": float(stack["density2d"][yi, xi]),
            }
        )
    return rows


def cells_3d_rows(stack):
    record = stack["record"]
    columns = ["X", "Y", "Z_original_um", "Z_corregida_um", "X_norm", "Y_norm", "AP_mm"]
    rows = []
    for _, cell in stack["valid"][columns].iterrows():
        row = {
            "Stack_ID": record["Stack_ID"], "Experimento": record["Experimento"],
            "Animal": record["Animal"], "Grupo": record["Grupo"], "Region": record["Region"],
            "Bregma_corte_mm": record["Bregma_mm"],
        }
        row.update({column: float(cell[column]) for column in columns})
        rows.append(row)
    return rows


def eligible_for_group(stack):
    record = stack["record"]
    return text(record["QC"]).upper() == "OK" and truthy(record["Usar_mapa_grupal"])


def is_vehicle_group(value):
    return normalized_label(value) in {"veh", "vehiculo", "vehicle", "control", "vv"}


def build_group_2d(stacks, output_dir):
    root = Path(output_dir)
    output_dir = root / "mapas_2D_grupo"
    difference_dir = root / "diferencias_2D"
    output_dir.mkdir(parents=True, exist_ok=True)
    difference_dir.mkdir(parents=True, exist_ok=True)
    buckets = defaultdict(list)
    for stack in stacks:
        if eligible_for_group(stack):
            record = stack["record"]
            buckets[(record["Grupo"], record["Region"])].append(stack)

    prepared = []
    for key, members in buckets.items():
        counts = np.sum([member["counts2d"] for member in members], axis=0)
        volume = np.sum([member["volume2d"] for member in members], axis=0)
        density = density_from_counts(counts, volume)
        prepared.append((key, members, counts, volume, density))

    summary, long_rows = [], []
    for (group, region), members, counts, volume, density in prepared:
        animals = {(m["record"]["Experimento"], m["record"]["Animal"]) for m in members}
        peak, index = peak_metrics(density)
        filename = output_dir / f"2D_normalizado_{safe_name(group)}_{region}.png"
        pdf = filename.with_suffix(".pdf")
        fig, ax = plt.subplots(figsize=(6.2, 5.4))
        image = ax.imshow(
            density, origin="lower", extent=(0, 1, 0, 1), cmap="magma",
            vmin=0, vmax=GROUP_2D_VMAX_MM3, interpolation="bilinear",
        )
        ax.set(
            title=f"Distribución 2D normalizada — {group} — {region}",
            xlabel=X_AXIS_LABEL, ylabel=Y_AXIS_LABEL,
            xlim=(0, 1), ylim=(0, 1), aspect="equal",
        )
        add_fixed_density_colorbar(fig, image, ax, GROUP_2D_VMAX_MM3)
        save_figure_atomic(fig, filename, dpi=220, bbox_inches="tight")
        save_figure_atomic(fig, pdf, bbox_inches="tight")
        plt.close(fig)
        bregmas = [float(m["record"]["Bregma_mm"]) for m in members]
        summary.append(
            {
                "Grupo": group, "Region": region,
                "N_animales": len(animals), "N_stacks": len(members), "Celulas": int(counts.sum()),
                "Volumen_um3": float(volume.sum()), "Densidad_global_mm3": float(counts.sum() / volume.sum() * 1e9),
                "Bregma_min_mm": min(bregmas), "Bregma_max_mm": max(bregmas),
                "Pico_local_mm3": peak, "Entropia": normalized_entropy(counts),
                "Escala_color_vmax_mm3": GROUP_2D_VMAX_MM3,
                "Archivo_PNG": str(filename), "Archivo_PDF": str(pdf),
            }
        )
        bins = counts.shape[0]
        for yi, xi in np.ndindex(counts.shape):
            long_rows.append(
                {
                    "Grupo": group, "Region": region,
                    "N_animales": len(animals), "N_stacks": len(members), "Bin_X": int(xi), "Bin_Y": int(yi),
                    "X_norm": (xi + 0.5) / bins, "Y_norm": (yi + 0.5) / bins,
                    "Celulas_sumadas": int(counts[yi, xi]), "Volumen_sumado_um3": float(volume[yi, xi]),
                    "Densidad_celulas_mm3": float(density[yi, xi]),
                }
            )
    differences = []
    prepared_differences = []
    by_region = defaultdict(list)
    for item in prepared:
        key = item[0]
        by_region[key[1]].append(item)
    for region, items in by_region.items():
        reference = next((item for item in items if is_vehicle_group(item[0][0])), None)
        if reference is None:
            continue
        ref_group, ref_density = reference[0][0], reference[4]
        for item in items:
            group, density = item[0][0], item[4]
            if group == ref_group:
                continue
            valid = np.isfinite(ref_density) & np.isfinite(density)
            difference = np.full_like(density, np.nan)
            difference[valid] = density[valid] - ref_density[valid]
            prepared_differences.append(
                (group, ref_group, region, density, ref_density, difference, valid)
            )

    difference_vmax = pooled_difference_vmax([item[5] for item in prepared_differences])

    for group, ref_group, region, density, ref_density, difference, valid in prepared_differences:
        filename = difference_dir / f"Dif_2D_{safe_name(group)}_menos_{safe_name(ref_group)}_{region}.png"
        pdf = filename.with_suffix(".pdf")
        fig, ax = plt.subplots(figsize=(6.3, 5.4))
        image = ax.imshow(
            difference, origin="lower", extent=(0, 1, 0, 1), cmap="coolwarm",
            vmin=-difference_vmax, vmax=difference_vmax, interpolation="bilinear",
        )
        ax.set(title=f"{group} − {ref_group} — {region}", xlabel=X_AXIS_LABEL, ylabel=Y_AXIS_LABEL, aspect="equal")
        fig.colorbar(image, ax=ax, label="Δ células/mm³")
        save_figure_atomic(fig, filename, dpi=220, bbox_inches="tight")
        save_figure_atomic(fig, pdf, bbox_inches="tight")
        plt.close(fig)
        bins = density.shape[0]
        for yi, xi in np.argwhere(valid):
            differences.append(
                {
                    "Comparacion": group, "Referencia": ref_group, "Region": region,
                    "Bin_X": int(xi), "Bin_Y": int(yi), "X_norm": (xi + 0.5) / bins,
                    "Y_norm": (yi + 0.5) / bins, "Densidad_comparacion_mm3": float(density[yi, xi]),
                    "Densidad_referencia_mm3": float(ref_density[yi, xi]),
                    "Diferencia_mm3": float(difference[yi, xi]),
                    "Archivo_PNG": str(filename), "Archivo_PDF": str(pdf),
                }
            )
    return pd.DataFrame(summary), pd.DataFrame(long_rows), pd.DataFrame(differences)


def global_3d_arrays(stack, xy_bins=BINS_XY_3D):
    region = stack["record"]["Region"]
    anterior, posterior = RANGES[region]
    # El margen contiene la profundidad real de stacks centrados en los límites
    # del atlas. Evita perder células o volumen sin extender artificialmente Z.
    ap_edges = np.arange(
        posterior - AP_MARGIN_MM,
        anterior + AP_MARGIN_MM + AP_VOXEL_MM * 1.001,
        AP_VOXEL_MM,
    )
    valid = stack["valid"]
    ap_values = valid["AP_mm"].to_numpy(float)
    if (ap_values < ap_edges[0]).any() or (ap_values > ap_edges[-1]).any():
        raise ValueError(
            "La profundidad del stack excede el margen AP configurado; "
            "no se generó el 3D para evitar perder células silenciosamente."
        )
    counts, _ = np.histogramdd(
        valid[["X_norm", "Y_norm", "AP_mm"]].to_numpy(),
        bins=(np.linspace(0, 1, xy_bins + 1), np.linspace(0, 1, xy_bins + 1), ap_edges),
    )

    xy_weights = np.full((xy_bins, xy_bins), 1 / (xy_bins * xy_bins))
    thickness_mm = stack["thickness_um"] / 1000
    bregma = float(stack["record"]["Bregma_mm"])
    interval = sorted([bregma - thickness_mm / 2, bregma + thickness_mm / 2])
    if interval[0] < ap_edges[0] or interval[1] > ap_edges[-1]:
        raise ValueError(
            "El volumen del stack excede el margen AP configurado; "
            "no se generó el 3D para evitar truncar volumen silenciosamente."
        )
    overlap = np.maximum(0, np.minimum(ap_edges[1:], interval[1]) - np.maximum(ap_edges[:-1], interval[0]))
    interval_width = max(interval[1] - interval[0], np.finfo(float).eps)
    z_weights = overlap / interval_width
    volume = float(stack["record"]["Volumen_um3"]) * xy_weights[:, :, None] * z_weights[None, None, :]
    return counts, volume, ap_edges


def plot_group_3d(group, region, counts, volume, ap_edges, filename, vmax=None):
    if vmax is None:
        vmax = GROUP_3D_VMAX_MM3
    density = density_from_counts(counts, volume)
    indices = np.argwhere((counts > 0) & np.isfinite(density))
    fig = plt.figure(figsize=(7.2, 6.2))
    ax3d = fig.add_subplot(111, projection="3d")
    if len(indices):
        values = density[tuple(indices.T)]
        scatter = ax3d.scatter(
            (indices[:, 0] + 0.5) / counts.shape[0],
            (indices[:, 1] + 0.5) / counts.shape[1],
            (ap_edges[indices[:, 2]] + ap_edges[indices[:, 2] + 1]) / 2,
            c=values, s=10 + 3 * counts[tuple(indices.T)], cmap="magma", vmin=0, vmax=vmax, alpha=0.75,
        )
        add_fixed_density_colorbar(
            fig, scatter, ax3d, vmax, shrink=0.66, pad=0.1,
        )
    ax3d.set(
        xlabel=X_AXIS_LABEL, ylabel=Y_AXIS_LABEL,
        zlabel="Bregma (mm)", title=f"Distribución espacial 3D — {group} — {region}",
    )
    ax3d.view_init(elev=23, azim=-55)
    save_figure_atomic(fig, filename, dpi=220, bbox_inches="tight")
    save_figure_atomic(fig, Path(filename).with_suffix(".pdf"), bbox_inches="tight")
    plt.close(fig)


def plot_difference_3d(group, reference, region, difference, ap_edges, filename, vmax=None):
    valid = np.isfinite(difference)
    if vmax is None:
        finite = np.abs(difference[valid])
        vmax = float(finite.max()) if len(finite) else 1.0
    vmax = max(float(vmax), np.finfo(float).eps)
    indices = np.argwhere(valid & (difference != 0))
    fig = plt.figure(figsize=(7.2, 6.2))
    ax3d = fig.add_subplot(111, projection="3d")
    if len(indices):
        values = difference[tuple(indices.T)]
        scatter = ax3d.scatter(
            (indices[:, 0] + 0.5) / difference.shape[0],
            (indices[:, 1] + 0.5) / difference.shape[1],
            (ap_edges[indices[:, 2]] + ap_edges[indices[:, 2] + 1]) / 2,
            c=values, s=13, cmap="coolwarm", vmin=-vmax, vmax=vmax, alpha=0.75,
        )
        fig.colorbar(scatter, ax=ax3d, shrink=0.62, pad=0.16, label="Δ células/mm³")
    ax3d.set(xlabel=X_AXIS_LABEL, ylabel=Y_AXIS_LABEL, zlabel="Bregma (mm)", title=f"{group} − {reference} — {region}")
    ax3d.view_init(elev=23, azim=-55)
    save_figure_atomic(fig, filename, dpi=220, bbox_inches="tight")
    save_figure_atomic(fig, Path(filename).with_suffix(".pdf"), bbox_inches="tight")
    plt.close(fig)


def build_group_3d(stacks, output_dir):
    root = Path(output_dir)
    map_dir, matrix_dir, difference_dir = root / "mapas_3D_grupo", root / "matrices_3D", root / "diferencias_3D"
    map_dir.mkdir(parents=True, exist_ok=True)
    matrix_dir.mkdir(parents=True, exist_ok=True)
    difference_dir.mkdir(parents=True, exist_ok=True)
    buckets = defaultdict(list)
    for stack in stacks:
        if eligible_for_group(stack):
            record = stack["record"]
            buckets[(record["Grupo"], record["Region"])].append(stack)

    prepared = []
    for (group, region), members in buckets.items():
        arrays = [global_3d_arrays(member) for member in members]
        counts = np.sum([item[0] for item in arrays], axis=0)
        volume = np.sum([item[1] for item in arrays], axis=0)
        ap_edges = arrays[0][2]
        density = density_from_counts(counts, volume)
        prepared.append((group, region, members, counts, volume, density, ap_edges))

    summaries, voxels = [], []
    for group, region, members, counts, volume, density, ap_edges in prepared:
        animals = {(m["record"]["Experimento"], m["record"]["Animal"]) for m in members}
        peak, index = peak_metrics(density)
        base = f"3D_{safe_name(group)}_{region}"
        figure = map_dir / f"{base}.png"
        matrix = matrix_dir / f"{base}.npz"
        plot_group_3d(
            group, region, counts, volume, ap_edges, figure,
            vmax=GROUP_3D_VMAX_MM3,
        )
        np.savez_compressed(
            matrix, conteos=counts, volumen_um3=volume, densidad_celulas_mm3=density,
            x_edges=np.linspace(0, 1, counts.shape[0] + 1),
            y_edges=np.linspace(0, 1, counts.shape[1] + 1),
            ap_edges_mm=ap_edges,
            convencion_x=X_COORDINATE_CONVENTION,
            convencion_y=Y_COORDINATE_CONVENTION,
        )
        row = {
            "Grupo": group, "Region": region, "N_animales": len(animals), "N_stacks": len(members),
            "Celulas": int(counts.sum()), "Volumen_muestreado_um3": float(volume.sum()),
            "Densidad_global_mm3": float(counts.sum() / volume.sum() * 1e9), "Pico_local_mm3": peak,
            "Entropia_3D": normalized_entropy(counts),
            "Escala_color_vmax_mm3": GROUP_3D_VMAX_MM3,
            "Archivo_mapa": str(figure), "Archivo_matriz": str(matrix),
        }
        if index is not None:
            row.update(
                {
                    "Hotspot_X_norm": (index[0] + 0.5) / counts.shape[0],
                    "Hotspot_Y_norm": (index[1] + 0.5) / counts.shape[1],
                    "Hotspot_AP_mm": float((ap_edges[index[2]] + ap_edges[index[2] + 1]) / 2),
                }
            )
        summaries.append(row)
        for xi, yi, zi in np.argwhere(volume > 0):
            ap_value = (ap_edges[zi] + ap_edges[zi + 1]) / 2
            voxels.append(
                {
                    "Grupo": group, "Region": region, "N_animales": len(animals), "N_stacks": len(members),
                    "Voxel_X": int(xi), "Voxel_Y": int(yi), "Voxel_AP": int(zi),
                    "X_norm": (xi + 0.5) / counts.shape[0], "Y_norm": (yi + 0.5) / counts.shape[1],
                    "AP_mm": float(ap_value),
                    "Celulas_sumadas": int(counts[xi, yi, zi]), "Volumen_sumado_um3": float(volume[xi, yi, zi]),
                    "Densidad_celulas_mm3": float(density[xi, yi, zi]),
                }
            )
    differences = []
    prepared_differences = []
    by_region = defaultdict(list)
    for item in prepared:
        by_region[item[1]].append(item)
    for region, items in by_region.items():
        reference = next((item for item in items if is_vehicle_group(item[0])), None)
        if reference is None:
            continue
        ref_group, ref_density = reference[0], reference[5]
        for group, _, members, counts, volume, density, ap_edges in items:
            if group == ref_group:
                continue
            valid = np.isfinite(ref_density) & np.isfinite(density)
            difference = np.full_like(density, np.nan)
            difference[valid] = density[valid] - ref_density[valid]
            prepared_differences.append(
                (group, ref_group, region, density, ref_density, difference, valid, ap_edges)
            )

    difference_vmax = pooled_difference_vmax([item[5] for item in prepared_differences])

    for group, ref_group, region, density, ref_density, difference, valid, ap_edges in prepared_differences:
        base = f"Dif_3D_{safe_name(group)}_menos_{safe_name(ref_group)}_{region}"
        figure = difference_dir / f"{base}.png"
        matrix = difference_dir / f"{base}.npz"
        plot_difference_3d(
            group, ref_group, region, difference, ap_edges, figure,
            vmax=difference_vmax,
        )
        np.savez_compressed(
            matrix,
            diferencia_celulas_mm3=difference,
            convencion_x=X_COORDINATE_CONVENTION,
            convencion_y=Y_COORDINATE_CONVENTION,
        )
        for xi, yi, zi in np.argwhere(valid):
            ap_value = (ap_edges[zi] + ap_edges[zi + 1]) / 2
            differences.append(
                {
                    "Comparacion": group, "Referencia": ref_group, "Region": region,
                    "Voxel_X": int(xi), "Voxel_Y": int(yi), "Voxel_AP": int(zi),
                    "X_norm": (xi + 0.5) / density.shape[0], "Y_norm": (yi + 0.5) / density.shape[1],
                    "AP_mm": float(ap_value),
                    "Densidad_comparacion_mm3": float(density[xi, yi, zi]),
                    "Densidad_referencia_mm3": float(ref_density[xi, yi, zi]),
                    "Diferencia_mm3": float(difference[xi, yi, zi]),
                    "Archivo_mapa": str(figure), "Archivo_matriz": str(matrix),
                }
            )
    return pd.DataFrame(summaries), pd.DataFrame(voxels), pd.DataFrame(differences)


def animal_summaries(stacks):
    buckets = defaultdict(list)
    for stack in stacks:
        if text(stack["record"]["QC"]).upper() == "OK":
            record = stack["record"]
            buckets[(record["Experimento"], record["Animal"], record["Grupo"], record["Region"])].append(stack)
    rows = []
    for (experiment, animal, group, region), members in buckets.items():
        cells = pd.concat([member["valid"] for member in members], ignore_index=True)
        arrays = [global_3d_arrays(member) for member in members]
        counts = np.sum([item[0] for item in arrays], axis=0)
        volume_grid = np.sum([item[1] for item in arrays], axis=0)
        ap_edges = arrays[0][2]
        density = density_from_counts(counts, volume_grid)
        peak, index = peak_metrics(density)
        total_volume = sum(float(member["record"]["Volumen_um3"]) for member in members)
        row = {
            "Experimento": experiment, "Animal": animal, "Grupo": group, "Region": region,
            "N_stacks": len(members), "Celulas_cFos": len(cells), "Volumen_um3": total_volume,
            "Densidad_celulas_mm3": len(cells) / total_volume * 1e9,
            "Densidad_x10_3_mm3": len(cells) / total_volume * 1e6,
            "Centroide_X_norm": float(cells["X_norm"].mean()), "Centroide_Y_norm": float(cells["Y_norm"].mean()),
            "Centroide_AP_mm": float(cells["AP_mm"].mean()),
            "SD_X_norm": float(cells["X_norm"].std(ddof=0)), "SD_Y_norm": float(cells["Y_norm"].std(ddof=0)),
            "SD_AP_mm": float(cells["AP_mm"].std(ddof=0)), "Pico_local_3D_mm3": peak,
            "Entropia_3D": normalized_entropy(counts),
        }
        if index is not None:
            row.update(
                {
                    "Hotspot_X_norm": (index[0] + 0.5) / counts.shape[0],
                    "Hotspot_Y_norm": (index[1] + 0.5) / counts.shape[1],
                    "Hotspot_AP_mm": float((ap_edges[index[2]] + ap_edges[index[2] + 1]) / 2),
                }
            )
        rows.append(row)
    return pd.DataFrame(rows)


def bregma_summaries(stacks):
    buckets = defaultdict(list)
    for stack in stacks:
        if text(stack["record"]["QC"]).upper() == "OK":
            r = stack["record"]
            key = (r["Experimento"], r["Animal"], r["Grupo"], r["Region"], float(r["Bregma_mm"]))
            buckets[key].append(stack)
    rows = []
    for (experiment, animal, group, region, bregma), members in buckets.items():
        cells = sum(len(member["valid"]) for member in members)
        volume = sum(float(member["record"]["Volumen_um3"]) for member in members)
        rows.append({
            "Experimento": experiment, "Animal": animal, "Grupo": group, "Region": region,
            "Bregma_mm": bregma, "N_stacks": len(members), "Celulas_cFos": cells,
            "Volumen_um3": volume, "Densidad_celulas_mm3": cells / volume * 1e9,
            "Densidad_x10_3_mm3": cells / volume * 1e6,
        })
    return pd.DataFrame(rows)


def sector_summaries(stacks):
    rows = []
    for stack in stacks:
        if text(stack["record"]["QC"]).upper() != "OK":
            continue
        record, valid = stack["record"], stack["valid"]
        fractions = stack["roi_fraction"]
        ny, nx = fractions.shape
        sectors = [
            ("Medial-Dorsal", valid["X_norm"].lt(.5) & valid["Y_norm"].ge(.5), fractions[ny // 2:, :nx // 2]),
            ("Lateral-Dorsal", valid["X_norm"].ge(.5) & valid["Y_norm"].ge(.5), fractions[ny // 2:, nx // 2:]),
            ("Medial-Ventral", valid["X_norm"].lt(.5) & valid["Y_norm"].lt(.5), fractions[:ny // 2, :nx // 2]),
            ("Lateral-Ventral", valid["X_norm"].ge(.5) & valid["Y_norm"].lt(.5), fractions[:ny // 2, nx // 2:]),
        ]
        total_fraction = max(float(fractions.sum()), np.finfo(float).eps)
        for sector, mask, sector_fraction in sectors:
            volume = float(record["Volumen_um3"]) * float(sector_fraction.sum()) / total_fraction
            cells = int(mask.sum())
            rows.append({
                "Experimento": record["Experimento"], "Animal": record["Animal"],
                "Grupo": record["Grupo"], "Region": record["Region"], "Sector": sector,
                "Celulas_cFos": cells, "Volumen_um3": volume,
            })
    frame = pd.DataFrame(rows)
    if frame.empty:
        return frame
    grouped = frame.groupby(["Experimento", "Animal", "Grupo", "Region", "Sector"], as_index=False).agg(
        Celulas_cFos=("Celulas_cFos", "sum"), Volumen_um3=("Volumen_um3", "sum")
    )
    grouped["Densidad_celulas_mm3"] = grouped["Celulas_cFos"] / grouped["Volumen_um3"] * 1e9
    grouped["Densidad_x10_3_mm3"] = grouped["Densidad_celulas_mm3"] / 1000
    return grouped


def grouped_stats(frame, keys):
    if frame.empty:
        return pd.DataFrame()
    rows = []
    for values, members in frame.groupby(keys, dropna=False):
        values = values if isinstance(values, tuple) else (values,)
        density = members["Densidad_celulas_mm3"].dropna()
        n = len(density)
        sd = density.std(ddof=1) if n > 1 else np.nan
        row = dict(zip(keys, values))
        row.update({
            "N_animales": n, "Media_densidad_mm3": density.mean(), "SD_densidad_mm3": sd,
            "SEM_densidad_mm3": sd / np.sqrt(n) if n > 1 else np.nan,
            "Media_x10_3_mm3": density.mean() / 1000,
            "SEM_x10_3_mm3": sd / np.sqrt(n) / 1000 if n > 1 else np.nan,
        })
        rows.append(row)
    return pd.DataFrame(rows)


def prism_wide_density(animal_df):
    """Organiza la densidad por columnas, lista para pegar/importar en Prism."""
    columns = {}
    for region in ("BLAa", "BLAp"):
        for group in group_order(animal_df["Grupo"] if not animal_df.empty else []):
            values = animal_df.loc[
                animal_df["Region"].eq(region) & animal_df["Grupo"].eq(group)
            ].sort_values(["Experimento", "Animal"])["Densidad_x10_3_mm3"].reset_index(drop=True)
            columns[f"{group}_{region}"] = values
    return pd.DataFrame(columns)


def group_order(values):
    values = list(dict.fromkeys(map(text, values)))
    preferred = [name for name in ("Vehículo", "SCOP-HB", "LiCl") if name in values]
    return preferred + sorted(name for name in values if name and name not in preferred)


def group_color(group, index=0):
    fallback = ["#8F8F8F", "#4C95C7", "#E53935", "#7A5AF8", "#00A67D"]
    return GROUP_COLORS.get(group, fallback[index % len(fallback)])


def _bar_points(ax, x, values, color, width=.68):
    values = np.asarray(values, dtype=float)
    values = values[np.isfinite(values)]
    if not len(values):
        return
    ax.bar(x, values.mean(), width=width, color=color, edgecolor="none", zorder=1)
    jitter = np.linspace(-width * .16, width * .16, len(values)) if len(values) > 1 else np.array([0.0])
    ax.scatter(x + jitter, values, s=30, color="#151515", zorder=3)
    if len(values) > 1:
        sem = values.std(ddof=1) / np.sqrt(len(values))
        ax.errorbar(x, values.mean(), yerr=sem, color="black", lw=2, capsize=5, zorder=4)


def plot_density_figures(animal_df, bregma_df, output_dir):
    main_dir = Path(output_dir) / "graficas_principales"
    supplement_dir = Path(output_dir) / "graficas_suplementarias"
    main_dir.mkdir(parents=True, exist_ok=True)
    supplement_dir.mkdir(parents=True, exist_ok=True)
    groups = group_order(animal_df["Grupo"] if not animal_df.empty else [])
    files = []
    for region in sorted(animal_df["Region"].unique()):
        subset = animal_df[animal_df["Region"].eq(region)]
        fig, ax = plt.subplots(figsize=(6.2, 5.8))
        for index, group in enumerate(groups):
            values = subset.loc[subset["Grupo"].eq(group), "Densidad_x10_3_mm3"]
            _bar_points(ax, index, values, group_color(group, index))
        ax.set(
            xticks=np.arange(len(groups)), xticklabels=groups,
            ylabel="c-Fos+ cells × 10³/mm³", title=f"Densidad c-Fos — {region}",
        )
        ax.spines[["top", "right"]].set_visible(False)
        fig.tight_layout()
        png = main_dir / f"Densidad_cFos_{region}.png"
        save_figure_atomic(fig, png, dpi=300, bbox_inches="tight")
        save_figure_atomic(fig, png.with_suffix(".pdf"), bbox_inches="tight")
        plt.close(fig)
        files.append(str(png))

    for region in sorted(bregma_df["Region"].unique() if not bregma_df.empty else []):
        subset = bregma_df[bregma_df["Region"].eq(region)]
        bregmas = sorted(subset["Bregma_mm"].dropna().unique())
        fig, ax = plt.subplots(figsize=(max(7, len(bregmas) * 1.45), 5.8))
        width = min(.24, .75 / max(1, len(groups)))
        offsets = (np.arange(len(groups)) - (len(groups) - 1) / 2) * width
        for gindex, group in enumerate(groups):
            for bindex, bregma in enumerate(bregmas):
                values = subset.loc[
                    subset["Grupo"].eq(group) & np.isclose(subset["Bregma_mm"], bregma),
                    "Densidad_x10_3_mm3",
                ]
                _bar_points(ax, bindex + offsets[gindex], values, group_color(group, gindex), width=width * .88)
        ax.set(
            xticks=np.arange(len(bregmas)), xticklabels=[f"{value:.2f}" for value in bregmas],
            xlabel="Bregma (mm)", ylabel="c-Fos+ cells × 10³/mm³",
            title=f"Densidad por bregma — {region}",
        )
        handles = [plt.Rectangle((0, 0), 1, 1, color=group_color(group, i)) for i, group in enumerate(groups)]
        ax.legend(handles, groups, frameon=False, title="Grupo")
        ax.spines[["top", "right"]].set_visible(False)
        fig.tight_layout()
        png = supplement_dir / f"Densidad_por_bregma_3_barras_{region}.png"
        save_figure_atomic(fig, png, dpi=300, bbox_inches="tight")
        save_figure_atomic(fig, png.with_suffix(".pdf"), bbox_inches="tight")
        plt.close(fig)
        files.append(str(png))
    return files


def style_results_workbook(path):
    workbook = load_workbook(path)
    try:
        header_fill = PatternFill("solid", fgColor="0B1F33")
        header_font = Font(color="FFFFFF", bold=True)
        for sheet in workbook.worksheets:
            sheet.sheet_view.showGridLines = False
            sheet.freeze_panes = "A2"
            if sheet.max_row:
                sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            sheet.row_dimensions[1].height = 32
            headers = {cell.value: cell.column for cell in sheet[1] if cell.value}
            for column in range(1, sheet.max_column + 1):
                values = [text(sheet.cell(row, column).value) for row in range(1, min(sheet.max_row, 200) + 1)]
                sheet.column_dimensions[get_column_letter(column)].width = min(38, max(10, max(map(len, values), default=0) + 2))
            for name, column in headers.items():
                fmt = "0.00" if "Bregma" in str(name) else "#,##0.0" if "Densidad" in str(name) else None
                if fmt:
                    for row in range(2, sheet.max_row + 1):
                        sheet.cell(row, column).number_format = fmt
        workbook.save(path)
    finally:
        workbook.close()


def run_final(project):
    project = Path(project)
    registry = project / "registro_general_stacks.xlsx"
    legacy = project / "metadata_stacks.xlsx"
    source = registry if registry.exists() else legacy
    if not source.exists():
        raise FileNotFoundError(
            "No existe registro_general_stacks.xlsx ni un metadata_stacks.xlsx legado."
        )
    metadata = pd.read_excel(source, sheet_name=0, dtype=str).fillna("")
    metadata = metadata.reindex(columns=META_COLUMNS, fill_value="")
    qc_ok = metadata["QC"].astype(str).str.strip().str.casefold().eq("ok")
    selected = metadata["Incluir_analisis"].map(selected_for_analysis)
    process = qc_ok & selected
    skipped = metadata.loc[~process, ["Stack_ID", "Experimento", "Animal", "QC", "Incluir_analisis", "Detalle_QC"]].copy()
    skipped["Motivo_exclusion"] = np.where(
        ~selected.loc[skipped.index],
        "Animal desmarcado para esta corrida",
        "QC distinto de OK",
    )
    metadata_to_process = metadata.loc[process].copy()
    if metadata_to_process.empty:
        raise RuntimeError("No hay stacks seleccionados con QC=OK para ejecutar el análisis final.")
    run_dir = project / "salidas" / datetime.now().strftime("corrida_%Y%m%d_%H%M%S")
    run_dir.mkdir(parents=True, exist_ok=True)

    stacks, stack_rows, errors, bins2d_rows, cells3d_rows = [], [], [], [], []
    for _, record in metadata_to_process.iterrows():
        try:
            resolved_record = resolve_record_file_paths(project, record.to_dict())
            stack = prepare_stack(resolved_record)
            stacks.append(stack)
            stack_rows.append(stack["record"])
            bins2d_rows.extend(stack_2d_rows(stack))
            cells3d_rows.extend(cells_3d_rows(stack))
            plot_stack_outputs(stack, run_dir)
        except Exception as exc:
            errors.append({"Stack_ID": record.get("Stack_ID", ""), "Error": str(exc)})

    if not stacks:
        details = "\n".join(f"{row['Stack_ID']}: {row['Error']}" for row in errors)
        raise RuntimeError("No se pudo procesar ningún stack.\n" + details)

    stack_df = pd.DataFrame(stack_rows)
    animal_df = animal_summaries(stacks)
    bregma_df = bregma_summaries(stacks)
    sector_df = sector_summaries(stacks)
    animal_stats_df = grouped_stats(animal_df, ["Grupo", "Region"])
    bregma_stats_df = grouped_stats(bregma_df, ["Grupo", "Region", "Bregma_mm"])
    sector_stats_df = grouped_stats(sector_df, ["Grupo", "Region", "Sector"])
    prism_audit_df = animal_df.pivot_table(
        index=["Experimento", "Animal", "Grupo"], columns="Region",
        values="Densidad_celulas_mm3", aggfunc="first",
    ).reset_index()
    prism_df = prism_wide_density(animal_df)
    group2d_df, group2d_bins_df, difference2d_df = build_group_2d(stacks, run_dir)
    group3d_df, group3d_voxels_df, difference3d_df = build_group_3d(stacks, run_dir)
    plot_density_figures(animal_df, bregma_df, run_dir)

    output_xlsx = run_dir / "resultados_cfos_2D_3D.xlsx"
    temporary_xlsx = run_dir / ".resultados_cfos_2D_3D.partial.xlsx"
    coordinate_dictionary = pd.DataFrame(
        [
            {
                "Coordenada": "X_ML_norm",
                "Convencion": X_COORDINATE_CONVENTION,
                "Transformacion": "X corregida por montaje y normalizada dentro de la ROI",
            },
            {
                "Coordenada": "Y_DV_norm",
                "Convencion": Y_COORDINATE_CONVENTION,
                "Transformacion": Y_COORDINATE_TRANSFORM,
            },
            {
                "Coordenada": "AP_mm",
                "Convencion": "Bregma real en mm",
                "Transformacion": "Z óptica corregida y centrada en el bregma del corte",
            },
        ]
    )
    with pd.ExcelWriter(temporary_xlsx, engine="openpyxl") as writer:
        pd.DataFrame(
            [{
                "Version_pipeline": APP_VERSION,
                "Fecha_corrida": datetime.now().isoformat(timespec="seconds"),
                "Convencion_X": X_COORDINATE_CONVENTION,
                "Convencion_Y": Y_COORDINATE_CONVENTION,
                "Transformacion_Y": Y_COORDINATE_TRANSFORM,
                "Escala_2D_grupal_vmax_celulas_mm3": GROUP_2D_VMAX_MM3,
                "Escala_3D_grupal_vmax_celulas_mm3": GROUP_3D_VMAX_MM3,
                "Escala_diferencias": "Simétrica; percentil 99.5 absoluto no nulo agrupado por modalidad",
            }]
        ).to_excel(writer, index=False, sheet_name="Version")
        coordinate_dictionary.to_excel(writer, index=False, sheet_name="Convenciones_espaciales")
        metadata.to_excel(writer, index=False, sheet_name="Stacks_registrados")
        skipped.to_excel(writer, index=False, sheet_name="Stacks_no_ejecutados")
        spatial_export_frame(stack_df).to_excel(writer, index=False, sheet_name="QC_y_metricas_stack")
        pd.DataFrame(errors).to_excel(writer, index=False, sheet_name="Errores")
        spatial_export_frame(animal_df).to_excel(writer, index=False, sheet_name="Analisis_principal_animal")
        animal_stats_df.to_excel(writer, index=False, sheet_name="Resumen_densidad_animal")
        prism_audit_df.to_excel(writer, index=False, sheet_name="Prism_grupo_x_region")
        prism_df.to_excel(writer, index=False, sheet_name="Prism_densidad_x10_3")
        bregma_df.to_excel(writer, index=False, sheet_name="Densidad_por_bregma")
        bregma_stats_df.to_excel(writer, index=False, sheet_name="Resumen_por_bregma")
        sector_df.to_excel(writer, index=False, sheet_name="Sectores_por_animal")
        sector_stats_df.to_excel(writer, index=False, sheet_name="Resumen_sectores")
        spatial_export_frame(pd.DataFrame(bins2d_rows)).to_excel(writer, index=False, sheet_name="Bins_2D_stack")
        spatial_export_frame(group2d_df).to_excel(writer, index=False, sheet_name="Resumen_2D_grupo")
        spatial_export_frame(group2d_bins_df).to_excel(writer, index=False, sheet_name="Bins_2D_grupo")
        spatial_export_frame(difference2d_df).to_excel(writer, index=False, sheet_name="Diferencias_2D")
        spatial_export_frame(pd.DataFrame(cells3d_rows)).to_excel(writer, index=False, sheet_name="Celulas_3D")
        spatial_export_frame(group3d_df).to_excel(writer, index=False, sheet_name="Resumen_3D_grupo")
        spatial_export_frame(group3d_voxels_df).to_excel(writer, index=False, sheet_name="Voxeles_3D_grupo")
        spatial_export_frame(difference3d_df).to_excel(writer, index=False, sheet_name="Diferencias_3D")
    style_results_workbook(temporary_xlsx)
    temporary_xlsx.replace(output_xlsx)

    notes = f"""RESULTADOS DEL PIPELINE c-Fos BLA — HEATMAP V1.2

ANÁLISIS PRINCIPAL: la unidad experimental es el animal. Para cada animal y
región se calcula densidad = suma(células c-Fos válidas) / suma(volumen ROI).
Las figuras principales muestran puntos individuales, media y SEM.

SELECCIÓN: solamente entran a esta corrida los animales marcados en la columna
Analizar de la pestaña Proyecto guardado y análisis final. La decisión se aplica
a todos los stacks BLAa/BLAp del animal. Los animales desmarcados permanecen en
el registro y se documentan en la hoja Stacks_no_ejecutados.

MONTAJE: todos los tejidos son hemisferio izquierdo. Montado de lado derecho=No
aplica Flip X; Montado de lado derecho=Sí aplica Flip Z. No existe un Flip Y
dependiente del montaje. Para la salida anatómica, Y_DV_norm = 1 - Y_imagen_norm:
Y=0 es ventral y Y=1 es dorsal. Z se conserva en µm y AP se expresa como
bregma real, sin normalizar Z.

2D: X/Y se transforman a una cuadrícula rectangular 0–1. X=0 es medial y X=1
es lateral; Y=0 es ventral y Y=1 es dorsal. El mapa no dibuja la silueta
BLAa/BLAp y resume la distribución relativa de todos los bregmas de la región.
La densidad se calcula como suma(células)/suma(volumen transformado). Todos los
mapas grupales 2D usan la misma escala fija: 0–80,000 células/mm³; los valores
superiores se conservan en las tablas y se saturan solamente en el color.

3D: figura descriptiva única por grupo/región con X/Y normalizadas y bregma real.
Los vóxeles sin volumen observado son NaN y no se interpretan como cero. Todos
los mapas grupales 3D usan la misma escala fija: 0–300,000 células/mm³; los
valores superiores se conservan en las tablas y se saturan solamente en el color.

EXCEL: Prism_densidad_x10_3 contiene una columna por grupo/región en las mismas
unidades de las figuras principales. Analisis_principal_animal conserva la
identidad de cada animal para auditoría. Las columnas espaciales exportadas
indican explícitamente ML o DV y la convención 0–1 en el nombre; la hoja
Convenciones_espaciales documenta las transformaciones.

SUPLEMENTO: densidad por bregma con tres barras separadas por posición, una por
grupo, puntos individuales y SEM. Los sectores y mapas espaciales son análisis
secundarios/exploratorios; no constituyen significancia vóxel a vóxel.
"""
    (run_dir / "LEEME_resultados.txt").write_text(notes, encoding="utf-8")
    if errors:
        detail = "\n".join(f"{row['Stack_ID']}: {row['Error']}" for row in errors)
        (run_dir / "ADVERTENCIA_stacks_con_error.txt").write_text(
            "La corrida terminó con stacks que no pudieron procesarse.\n"
            "No se incluyeron en los resúmenes. Revisa y corrige antes de interpretar.\n\n"
            + detail,
            encoding="utf-8",
        )
    return run_dir
